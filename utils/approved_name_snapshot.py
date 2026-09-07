"""Build a private glossary snapshot from explicitly selected approved name rows."""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook


def load_snapshot(path: Path, target_langs: list[str]) -> list[dict]:
    data = json.loads(path.read_text(encoding="utf-8-sig"))
    if data.get("schema") != "approved-name-snapshot-v1" or not isinstance(data.get("terms"), list):
        raise ValueError("invalid approved name snapshot")
    seen = set()
    for term in data["terms"]:
        if not isinstance(term, dict) or not isinstance(term.get("source"), str) or not term["source"].strip():
            raise ValueError("snapshot requires nonempty term sources")
        if term["source"] in seen:
            raise ValueError("duplicate snapshot source")
        seen.add(term["source"])
        translations = term.get("translations")
        if not isinstance(translations, dict) or any(
            not isinstance(translations.get(lang), str) or not translations[lang].strip()
            for lang in target_langs
        ):
            raise ValueError("snapshot missing requested language")
    return data["terms"]


def build_snapshot(*, term_base: Path | None, names_workbook: Path, sheet_name: str,
                   name_rows: list[int], target_langs: list[str], output: Path) -> dict:
    from utils.large_text_multilingual_pack import _load_terms, _normalize_headers, _find_header, _language_columns
    from utils.language_config import SOURCE_HEADERS, target_header_candidates

    if not name_rows or len(name_rows) != len(set(name_rows)) or min(name_rows) < 2:
        raise ValueError("explicit unique name rows after the header are required")
    protected = [names_workbook] + ([term_base] if term_base else [])
    if output.resolve() in [p.resolve() for p in protected]:
        raise ValueError("snapshot must not overwrite input")
    langs = [lang.upper() for lang in target_langs]
    terms = {term["source"]: term for term in _load_terms(term_base, langs)}
    workbook = load_workbook(names_workbook, read_only=True, data_only=True)
    definitions = []
    try:
        sheet = workbook[sheet_name]
        if sheet.max_row == 1 and sheet.max_column == 1:
            sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        headers = _normalize_headers(next(iterator, ()))
        source_col = _find_header(headers, {h.lower() for h in SOURCE_HEADERS})
        english_col = _find_header(headers, target_header_candidates("en"))
        id_col = _find_header(headers, {"id", "key", "索引id"})
        columns = _language_columns(headers, langs, excluded_columns={id_col} if id_col is not None else set())
        if source_col is None:
            raise ValueError("missing source column in approved name definitions")
        wanted = set(name_rows)
        for number, values in enumerate(iterator, 2):
            if number not in wanted:
                continue
            wanted.remove(number)
            get = lambda col: str(values[col] or "").strip() if col is not None and col < len(values) else ""
            source = get(source_col)
            translations = {lang: get(col) for lang, col in columns.items()}
            if not source or not all(translations.values()):
                raise ValueError(f"incomplete approved name definition at row {number}")
            old = terms.get(source)
            if old and (any(old["translations"].get(lang) != translations[lang] for lang in langs)
                        or (old.get("reference_en") and get(english_col) and old["reference_en"] != get(english_col))):
                raise ValueError(f"approved name conflicts with glossary at row {number}; adjudicate before snapshot")
            origin = {"file": names_workbook.name, "sheet": sheet_name, "row": number, "id": get(id_col)}
            terms[source] = {**(old or {}), "source": source, "translations": translations,
                             "reference_en": get(english_col), "category": "person name",
                             "required": True, "strict": True, "approved_origin": origin}
            definitions.append(origin)
        if wanted:
            raise ValueError(f"approved name rows not found: {sorted(wanted)}")
    finally:
        workbook.close()
    result = {"schema": "approved-name-snapshot-v1", "target_languages": langs,
              "inputs": [{"file": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()} for p in protected],
              "approved_definitions": definitions,
              "terms": sorted(terms.values(), key=lambda term: len(term["source"]), reverse=True)}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--term-base", type=Path)
    parser.add_argument("--names-workbook", type=Path, required=True)
    parser.add_argument("--names-sheet", required=True)
    parser.add_argument("--name-rows", required=True, help="Explicit approved definition rows, comma-separated")
    parser.add_argument("--target-langs", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    result = build_snapshot(term_base=args.term_base, names_workbook=args.names_workbook,
                            sheet_name=args.names_sheet, name_rows=[int(x) for x in args.name_rows.split(",")],
                            target_langs=args.target_langs.split(","), output=args.output)
    print(json.dumps({"approved_names": len(result["approved_definitions"]), "terms": len(result["terms"])}))
    return 0
