"""Quality regression harness for localization outputs.

The harness is intentionally independent from one project workbook. It can run
small string fixtures and scan Excel language tables so quality regressions are
caught before final delivery.
"""
from __future__ import annotations

import html
import json
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from utils.readability_checker import check_readability
from utils.term_checker import check_chinese_residue
from utils.variable_checker import CheckResult, check_all as check_variables

HTML_ENTITY_PATTERN = re.compile(r'&(?:#[0-9]+|#x[0-9A-Fa-f]+|[A-Za-z][A-Za-z0-9]+);')
INTERNAL_TOKEN_PATTERN = re.compile(r'\b[A-Z]{2,}[A-Z0-9]*\d[A-Z0-9]*\b')
HASH_CODE_PATTERN = re.compile(r'#[A-Z]{2,}(?:##\d+|#\d+)*\b')
LETTER_PLACEHOLDER_COMPACTION_PATTERN = re.compile(r'\b[A-Z]{1,6}(?:##\d+|#\d+){2,}[A-Z0-9#]*\b')
PLACEHOLDER_WORD_GLUE_PATTERN = re.compile(r'##\d+[A-Za-z]{2,}##\d+')
ORPHAN_LEADING_CLITIC_PATTERN = re.compile(r"^\s*['’]s\b", re.IGNORECASE)
FULLWIDTH_PUNCTUATION_PATTERN = re.compile(r'[，。！？：；（）【】％＋－]')
WORD_START_PATTERN = re.compile(r'[A-Za-z]')

DEFAULT_HARD_ISSUES = {
    'variable_missing',
    'variable_extra',
    'variable_order',
    'bbcode_open_mismatch',
    'bbcode_close_mismatch',
    'bbcode_unclosed',
    'bbcode_color_mismatch',
    'newline_mismatch',
    'chinese_residue',
    'opaque_abbreviation',
    'clipped_word',
    'title_case_overuse',
    'internal_token_leak',
    'hash_code_abbreviation',
    'placeholder_compaction',
    'placeholder_word_glue',
    'html_entity_leak',
    'orphan_leading_clitic',
    'leading_lowercase',
    'punctuation_corruption',
    'fullwidth_punctuation',
}


@dataclass
class HarnessResult:
    passed: bool
    total_cases: int = 0
    rows_scanned: int = 0
    issue_counts: Counter = field(default_factory=Counter)
    issues: list[dict] = field(default_factory=list)
    failures: list[dict] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            'passed': self.passed,
            'total_cases': self.total_cases,
            'rows_scanned': self.rows_scanned,
            'issue_counts': dict(self.issue_counts),
            'issues': self.issues,
            'failures': self.failures,
        }


def load_fixture(path: str | Path) -> dict:
    return json.loads(Path(path).read_text(encoding='utf-8'))


def check_row(row_id, source: str, translation: str, lang: str = 'en') -> list[CheckResult]:
    """Run all row-level hard gates used by the harness."""
    results: list[CheckResult] = []
    source = str(source or '')
    translation = str(translation or '')

    results.extend(check_variables(row_id, source, translation))
    results.extend(check_chinese_residue(row_id, translation))
    results.extend(check_readability(row_id, source, translation, lang=lang))
    results.extend(_check_surface_regressions(row_id, source, translation))

    if any(r.check_type == 'internal_token_leak' for r in results):
        results = [r for r in results if r.check_type != 'opaque_abbreviation']

    return results


def run_fixture(fixture: dict, lang: str = 'en') -> HarnessResult:
    cases = fixture.get('cases', [])
    result = HarnessResult(passed=True, total_cases=len(cases))

    for case in cases:
        row_id = case.get('id', '')
        case_lang = case.get('lang', lang)
        issues = check_row(
            row_id=row_id,
            source=case.get('source', ''),
            translation=case.get('translation', ''),
            lang=case_lang,
        )
        actual = sorted({issue.check_type for issue in issues})
        expected = sorted(case.get('expected_issues', []))

        for issue_type in actual:
            result.issue_counts[issue_type] += 1

        if actual != expected:
            result.passed = False
            result.failures.append({
                'id': row_id,
                'source': case.get('source', ''),
                'translation': case.get('translation', ''),
                'expected_issues': expected,
                'actual_issues': actual,
            })

    return result


def scan_workbook(
    path: str | Path,
    lang: str = 'en',
    fail_on: Iterable[str] | None = None,
) -> HarnessResult:
    """Scan a workbook language table.

    The scanner expects either headers containing ID/CN/EN-like names or a
    simple first-three-column layout: ID, source, target.
    """
    fail_set = set(fail_on or DEFAULT_HARD_ISSUES)
    result = HarnessResult(passed=True)
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, read_only=True, data_only=False)

    try:
        for ws in wb.worksheets:
            id_col, src_col, tgt_col = _detect_columns(ws)
            if src_col is None or tgt_col is None:
                continue
            max_col = max(c for c in (id_col, src_col, tgt_col) if c is not None) + 1
            for row_index, row in enumerate(
                ws.iter_rows(min_row=2, max_col=max_col, values_only=True),
                start=2,
            ):
                row_id = row[id_col] if id_col is not None else row_index
                source = row[src_col]
                target = row[tgt_col]
                if not isinstance(source, str) or not isinstance(target, str):
                    continue

                result.rows_scanned += 1
                for issue in check_row(row_id, source, target, lang=lang):
                    result.issue_counts[issue.check_type] += 1
                    if issue.check_type in fail_set:
                        result.passed = False
                        result.issues.append({
                            'file': str(workbook_path),
                            'sheet': ws.title,
                            'row': row_index,
                            'id': row_id,
                            'check_type': issue.check_type,
                            'severity': issue.severity,
                            'message': issue.message,
                            'source': source,
                            'translation': target,
                            'auto_fix': issue.auto_fix,
                        })
    finally:
        wb.close()

    return result


def merge_results(results: list[HarnessResult]) -> HarnessResult:
    merged = HarnessResult(passed=all(r.passed for r in results))
    for result in results:
        merged.total_cases += result.total_cases
        merged.rows_scanned += result.rows_scanned
        merged.issue_counts.update(result.issue_counts)
        merged.issues.extend(result.issues)
        merged.failures.extend(result.failures)
    return merged


def _check_surface_regressions(row_id, source: str, translation: str) -> list[CheckResult]:
    results: list[CheckResult] = []

    if HTML_ENTITY_PATTERN.search(translation):
        results.append(_issue(
            row_id,
            'html_entity_leak',
            'HTML entity leaked into translation',
            source,
            translation,
            auto_fix=html.unescape(translation),
        ))

    token_match = INTERNAL_TOKEN_PATTERN.search(translation)
    if token_match and not _looks_like_allowed_runtime_code(token_match.group(0), source):
        results.append(_issue(
            row_id,
            'internal_token_leak',
            f"Internal token-like text leaked: {token_match.group(0)}",
            source,
            translation,
        ))

    hash_match = HASH_CODE_PATTERN.search(translation)
    if hash_match:
        results.append(_issue(
            row_id,
            'hash_code_abbreviation',
            f"Hash-prefixed code abbreviation leaked: {hash_match.group(0)}",
            source,
            translation,
        ))

    compact_match = LETTER_PLACEHOLDER_COMPACTION_PATTERN.search(translation)
    if compact_match:
        results.append(_issue(
            row_id,
            'placeholder_compaction',
            f"Letters are compacted into placeholders: {compact_match.group(0)}",
            source,
            translation,
        ))

    glued_match = PLACEHOLDER_WORD_GLUE_PATTERN.search(translation)
    if glued_match:
        results.append(_issue(
            row_id,
            'placeholder_word_glue',
            f"Word is glued between placeholders: {glued_match.group(0)}",
            source,
            translation,
        ))

    if ORPHAN_LEADING_CLITIC_PATTERN.search(translation):
        results.append(_issue(
            row_id,
            'orphan_leading_clitic',
            "Translation starts with orphan possessive clitic",
            source,
            translation,
        ))

    if _has_leading_lowercase(source, translation):
        results.append(_issue(
            row_id,
            'leading_lowercase',
            "Sentence-like translation starts with lowercase",
            source,
            translation,
            severity='warning',
        ))

    if _has_punctuation_corruption(source, translation):
        results.append(_issue(
            row_id,
            'punctuation_corruption',
            "Suspicious quote or question punctuation corruption",
            source,
            translation,
        ))

    if FULLWIDTH_PUNCTUATION_PATTERN.search(translation):
        results.append(_issue(
            row_id,
            'fullwidth_punctuation',
            "Fullwidth punctuation remains in Latin-script translation",
            source,
            translation,
        ))

    return results


def _detect_columns(ws) -> tuple[int | None, int | None, int | None]:
    first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(v or '').strip().lower() for v in first]

    def pick(candidates: set[str], fallback: int | None) -> int | None:
        for idx, header in enumerate(headers):
            if header in candidates:
                return idx
        return fallback if fallback is None or fallback < len(headers) else None

    id_col = pick({'id', 'key'}, 0)
    src_col = pick({'cn', 'zh', '中文', '原文', 'source', 'original'}, 1)
    tgt_col = pick({'en', 'english', '译文', 'translation', 'target'}, 2)
    return id_col, src_col, tgt_col


def _issue(
    row_id,
    check_type: str,
    message: str,
    source: str,
    translation: str,
    severity: str = 'error',
    auto_fix: str = '',
) -> CheckResult:
    return CheckResult(
        row_id=row_id,
        check_type=check_type,
        severity=severity,
        message=message,
        original=source,
        translation=translation,
        auto_fix=auto_fix,
    )


def _looks_like_allowed_runtime_code(token: str, source: str) -> bool:
    if token in source:
        return True
    # Common version-style or short gameplay terms are handled elsewhere.
    return bool(re.fullmatch(r'(?:HP|ATK|DEF|DMG|DPS|PVP|PVE|VIP|FPS|SFX|UI|ID)\d*', token))


def _visible_start(text: str) -> str:
    stripped = str(text)
    token_pattern = re.compile(r'^\s*(?:\\n|\n|<[^>]+>|\{[^}]+\}|##\d+|\[(?:[A-Za-z]+\d+|\d+)\])+')
    while True:
        new = token_pattern.sub('', stripped)
        if new == stripped:
            break
        stripped = new
    stripped = re.sub(r'^\s*(?:\\n|\n|\d+[\.)]\s*)+', '', stripped)
    return stripped.lstrip()


def _has_leading_lowercase(source: str, translation: str) -> bool:
    if re.match(r'^\s*\d', translation):
        return False
    if _starts_with_runtime_payload(translation):
        return False
    visible = _visible_start(translation)
    match = WORD_START_PATTERN.search(visible)
    if not match:
        return False
    char = match.group(0)
    if not char.islower():
        return False
    source_visible_len = len(re.sub(r'\s+', '', source))
    return source_visible_len >= 8 or bool(re.search(r'[。！？!?]$', source))


def _starts_with_runtime_payload(text: str) -> bool:
    return bool(re.match(
        r'^\s*(?:<[^>]+>\s*)*(?:##\d+|\{[^}]+\}|\[[A-Za-z]+\d+\])',
        str(text),
    ))


def _has_punctuation_corruption(source: str, translation: str) -> bool:
    if translation.count('"') % 2 == 1 and '"' not in source:
        return True
    source_asks_question = bool(re.search(r'[？?]|吗|么|什么|怎么|为何|是否', source))
    if source_asks_question and '?' not in translation and '"' in translation:
        return True
    return False
