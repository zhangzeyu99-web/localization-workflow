"""Bind reviewed glossary enrichment to exact translation inputs before API work."""
from __future__ import annotations

import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

from utils.language_config import normalize_language_code


def fingerprint(path: Path) -> dict:
    path = path.resolve()
    return {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def _langs(values: list[str]) -> list[str]:
    return sorted({normalize_language_code(value).upper() for value in values})


def _read(path: Path) -> dict:
    return json.loads(path.read_text(encoding='utf-8-sig'))


def _write(path: Path, data: dict, protected: list[Path]) -> None:
    if path.resolve() in {item.resolve() for item in protected}:
        raise ValueError('glossary output must not overwrite an input')
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def _check_assets(assets: list[dict]) -> None:
    for asset in assets:
        if fingerprint(Path(asset['path'])) != asset:
            raise ValueError(f"glossary stage input drift: {asset['path']}")


def _candidates(data: dict) -> dict[str, dict]:
    result = {}
    if not isinstance(data.get('candidates'), list):
        raise ValueError('glossary candidates array is required')
    for item in data['candidates'] + data.get('rejected', []):
        cn = str(item.get('cn') or '').strip()
        if not cn or cn in result:
            raise ValueError('empty or duplicate glossary candidate')
        result[cn] = item
    return result


def compact_review_packet(payload: dict) -> dict:
    """Keep full evidence on disk; send only actionable terms and short context."""
    def compact(item: dict) -> dict:
        cn = str(item['cn'])
        evidence = []
        seen = set()
        for source in item.get('evidence', []):
            text = str(source.get('source') or '')
            at = text.find(cn)
            start = max(0, at - 80) if at >= 0 else 0
            excerpt = text[start:start + max(240, len(cn) + 80)]
            if excerpt in seen:
                continue
            seen.add(excerpt)
            evidence.append({**{key: source[key] for key in ('file', 'sheet', 'row', 'id') if key in source},
                             'excerpt': excerpt, 'excerpt_only': excerpt != text})
            if len(evidence) == 2:
                break
        return {key: item[key] for key in ('cn', 'category', 'decision', 'reason', 'translations', 'missing_languages') if key in item} | {'evidence': evidence}

    candidates = list(_candidates(payload).values())
    return {
        'candidates': [compact(item) for item in candidates if item.get('decision') != 'reject'],
        'existing_incomplete': [compact(item) for item in payload.get('existing_incomplete', [])],
        'excluded_from_model_review': sum(item.get('decision') == 'reject' for item in candidates),
        'instructions': 'Judge durable naming value and translate only approved terms in one review. '
                        'Category, brevity or frequency alone is not approval. Keep ordinary prompts and '
                        'compositional variants in batch translation memory, not the permanent glossary. '
                        'Read full evidence by its source reference only when the excerpt is insufficient. '
                        'Do not scan or translate unrelated blank glossary rows.',
    }


def prepare_request(*, inputs: list[Path], term_base: Path, target_langs: list[str],
                    candidates: Path, output: Path) -> dict:
    payload = _read(candidates)
    data = {'schema': 'glossary-enrichment-request-v1', 'status': 'awaiting_review',
            'inputs': [fingerprint(path) for path in inputs], 'base': fingerprint(term_base),
            'candidates': fingerprint(candidates), 'target_languages': _langs(target_langs),
            'review_packet': compact_review_packet(payload)}
    _write(output, data, [*inputs, term_base, candidates])
    return data


def _workbook_rows(path: Path) -> dict:
    wb = load_workbook(path, read_only=True, data_only=False)
    try:
        result = {}
        for sheet in wb:
            if sheet.max_row == 1 and sheet.max_column == 1:
                sheet.reset_dimensions()
            rows = []
            for row in sheet.values:
                values = list(row)
                while values and values[-1] is None:
                    values.pop()
                rows.append(values)
            result[sheet.title] = rows
        return result
    finally:
        wb.close()


def accept_glossary(*, request: Path, reviewed: Path, updated: Path, output: Path) -> dict:
    data = _read(request)
    if data.get('schema') != 'glossary-enrichment-request-v1':
        raise ValueError('invalid glossary request')
    original_assets = [*data['inputs'], data['base'], data['candidates']]
    _check_assets(original_assets)
    raw_payload = _read(Path(data['candidates']['path']))
    raw = _candidates(raw_payload)
    decisions = _candidates(_read(reviewed))
    if set(raw) != set(decisions):
        raise ValueError('every extracted candidate requires an explicit review decision')
    included = {}
    for cn, item in decisions.items():
        decision = item.get('decision')
        if decision == 'include':
            if not item.get('category') or not item.get('evidence'):
                raise ValueError(f'candidate requires category and source evidence: {cn}')
            included[cn] = item
        elif decision != 'reject' or not item.get('reason'):
            raise ValueError(f'unresolved glossary candidate: {cn}')
    old = _workbook_rows(Path(data['base']['path']))
    new = _workbook_rows(updated)
    if list(old) != list(new):
        raise ValueError('glossary sheets must remain unchanged')
    appended = {}
    existing_targets = {}
    old_names = set()
    filled_existing_cells = 0
    for name, old_rows in old.items():
        rows = new[name]
        if len(rows) < len(old_rows):
            raise ValueError('existing glossary rows removed')
        if not old_rows:
            if rows:
                raise ValueError('cannot append terms to an untyped empty glossary sheet')
            continue
        headers = [str(x).strip().lower() for x in old_rows[0]]
        from utils.language_config import target_header_candidates
        cols = {}
        for lang in data['target_languages']:
            canonical = normalize_language_code(lang)
            options = target_header_candidates(canonical)
            # Prefer IDN to its ID alias, which may be the identifier column.
            cols[lang] = (headers.index(canonical) if canonical in headers
                          else next((i for i, h in enumerate(headers) if h in options), None))
        target_cols = {col for col in cols.values() if col is not None}
        for index, old_row in enumerate(old_rows):
            new_row = rows[index]
            for col in range(max(len(old_row), len(new_row))):
                before = old_row[col] if col < len(old_row) else None
                after = new_row[col] if col < len(new_row) else None
                if before == after:
                    continue
                if index > 0 and col in target_cols and before in (None, '') and isinstance(after, str) and after.strip():
                    filled_existing_cells += 1
                else:
                    raise ValueError('existing glossary rows changed or reordered')
        if 'cn' not in headers:
            if len(rows) != len(old_rows):
                raise ValueError('append requires glossary CN header')
            continue
        cn_col = headers.index('cn')
        old_names.update(str(r[cn_col]).strip() for r in old_rows[1:] if len(r) > cn_col and r[cn_col])
        for lang in data['target_languages']:
            if cols[lang] is None:
                raise ValueError(f'missing glossary language column: {lang}')
        category_col = next((i for i, h in enumerate(headers)
                             if h in {'分类', '类别', '术语分类', 'category', 'type'}), None)
        for row in rows[1:len(old_rows)]:
            cn = str(row[cn_col] or '').strip() if len(row) > cn_col else ''
            if cn:
                existing_targets[cn] = {lang: row[col] if len(row) > col else None
                                        for lang, col in cols.items()}
        for row in rows[len(old_rows):]:
            cn = str(row[cn_col] or '').strip() if len(row) > cn_col else ''
            if not cn or cn in appended or cn in old_names:
                raise ValueError('duplicate or empty appended glossary term')
            translations = {lang: row[col] if len(row) > col else None for lang, col in cols.items()}
            if any(not isinstance(value, str) or not value.strip() for value in translations.values()):
                raise ValueError(f'missing requested glossary translation: {cn}')
            if category_col is not None and cn in included:
                category = row[category_col] if len(row) > category_col else None
                if category != included[cn]['category']:
                    raise ValueError(f'approved glossary category changed: {cn}')
            appended[cn] = translations
    for item in raw_payload.get('existing_incomplete', []):
        translations = existing_targets.get(item['cn'], {})
        if any(not isinstance(translations.get(lang), str) or not translations[lang].strip()
               for lang in data['target_languages']):
            raise ValueError(f"missing requested existing glossary translation: {item['cn']}")
    if set(appended) != set(included):
        raise ValueError('appended glossary terms must match approved candidates exactly')
    for cn, translations in appended.items():
        for lang, expected in included[cn].get('translations', {}).items():
            key = normalize_language_code(lang).upper()
            if expected and key in translations and translations[key] != expected:
                raise ValueError(f'approved glossary translation changed: {cn}/{key}')
    assets = [*original_assets, fingerprint(request), fingerprint(reviewed), fingerprint(updated)]
    result = {'schema': 'glossary-enrichment-receipt-v1', 'status': 'accepted',
              'inputs': data['inputs'], 'updated': fingerprint(updated), 'assets': assets,
              'target_languages': data['target_languages'], 'approved_terms': len(included),
              'filled_existing_cells': filled_existing_cells,
              'rejected_terms': len(decisions) - len(included)}
    _write(output, result, [Path(item['path']) for item in assets])
    return result


def verify_receipt(receipt: Path, *, inputs: list[Path], term_base: Path | None,
                   target_langs: list[str]) -> dict:
    data = _read(receipt)
    if data.get('schema') != 'glossary-enrichment-receipt-v1' or data.get('status') != 'accepted':
        raise ValueError('glossary enrichment has not been accepted')
    if _langs(target_langs) != data['target_languages']:
        raise ValueError('glossary receipt target languages mismatch')
    _check_assets(data['assets'])
    if {str(p.resolve()) for p in inputs} != {a['path'] for a in data['inputs']}:
        raise ValueError('glossary receipt source inputs mismatch')
    if term_base is None or fingerprint(term_base) != data['updated']:
        raise ValueError('glossary receipt updated term base mismatch')
    return data
