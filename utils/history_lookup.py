"""Exact, read-only lookup of accepted workbook translations for small tasks."""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .language_config import SOURCE_HEADERS, SUPPORTED_TRANSLATION_LANGUAGES, normalize_language_code, target_header_candidates


_SHORT_TARGET_HEADERS = {
    "en": {"英"},
    "fr": {"法"},
    "de": {"德"},
    "es": {"西"},
    "pt": {"葡"},
    "tr": {"土"},
    "ru": {"俄"},
    "vi": {"越"},
    "th": {"泰"},
    "ko": {"韩"},
    "ja": {"日"},
    "it": {"意"},
    "ar": {"阿"},
    "idn": {"印"},
}


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower()


def normalize_source_text(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _target_headers(lang: str) -> set[str]:
    return target_header_candidates(lang) | {_normalize_header(value) for value in _SHORT_TARGET_HEADERS.get(lang, set())}


@dataclass(frozen=True)
class HeaderMap:
    row: int
    source_col: int
    target_cols: dict[str, int]


@dataclass
class HistoryMatch:
    source_text: str
    translations: dict[str, str] = field(default_factory=dict)
    evidence: list[dict[str, object]] = field(default_factory=list)

    def status(self, languages: Iterable[str]) -> str:
        requested = set(languages)
        available = set(self.translations)
        if requested.issubset(available):
            return "exact"
        if available:
            return "partial"
        return "miss"


def _find_header_map(worksheet, languages: tuple[str, ...], *, require_targets: bool) -> HeaderMap | None:
    source_headers = {_normalize_header(value) for value in SOURCE_HEADERS}
    for row_index, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row), values_only=True),
        start=1,
    ):
        normalized = [_normalize_header(value) for value in row]
        source_col = next((index + 1 for index, value in enumerate(normalized) if value in source_headers), None)
        if source_col is None:
            continue
        target_cols: dict[str, int] = {}
        for lang in languages:
            candidates = _target_headers(lang)
            target_col = next((index + 1 for index, value in enumerate(normalized) if value in candidates), None)
            if target_col is not None:
                target_cols[lang] = target_col
        if not require_targets or target_cols:
            return HeaderMap(row=row_index, source_col=source_col, target_cols=target_cols)
    return None


def extract_source_queries(input_path: Path) -> list[str]:
    """Read unique current-task source strings without loading formatting or drawings."""
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        queries: list[str] = []
        seen: set[str] = set()
        for worksheet in workbook.worksheets:
            header = _find_header_map(worksheet, (), require_targets=False)
            if header is None:
                continue
            for row in worksheet.iter_rows(min_row=header.row + 1, min_col=header.source_col, max_col=header.source_col, values_only=True):
                source_text = normalize_source_text(row[0])
                if source_text and source_text not in seen:
                    seen.add(source_text)
                    queries.append(source_text)
        return queries
    finally:
        workbook.close()


def lookup_exact_history(
    queries: Iterable[str],
    history_paths: Iterable[Path],
    languages: Iterable[str],
) -> dict[str, object]:
    """Stream exact source matches and selected target columns; never fuzzy-search or render."""
    normalized_languages = tuple(dict.fromkeys(normalize_language_code(lang) for lang in languages))
    unsupported = [lang for lang in normalized_languages if lang not in SUPPORTED_TRANSLATION_LANGUAGES]
    if unsupported:
        raise ValueError(f"unsupported languages: {', '.join(unsupported)}")
    normalized_queries = list(dict.fromkeys(filter(None, (normalize_source_text(query) for query in queries))))
    matches = {query: HistoryMatch(source_text=query) for query in normalized_queries}
    unresolved = set(normalized_queries)
    scanned_files: list[str] = []

    for history_path in history_paths:
        if not unresolved:
            break
        history_path = Path(history_path)
        workbook = load_workbook(history_path, read_only=True, data_only=True)
        scanned_files.append(str(history_path))
        try:
            for worksheet in workbook.worksheets:
                if not unresolved:
                    break
                header = _find_header_map(worksheet, normalized_languages, require_targets=True)
                if header is None:
                    continue
                max_col = max(header.source_col, *header.target_cols.values())
                for excel_row, row in enumerate(
                    worksheet.iter_rows(min_row=header.row + 1, min_col=1, max_col=max_col, values_only=True),
                    start=header.row + 1,
                ):
                    source_text = normalize_source_text(row[header.source_col - 1])
                    if source_text not in unresolved:
                        continue
                    match = matches[source_text]
                    added: list[str] = []
                    for lang, target_col in header.target_cols.items():
                        translation = normalize_source_text(row[target_col - 1])
                        if translation and lang not in match.translations:
                            match.translations[lang] = translation
                            added.append(lang)
                    if added:
                        match.evidence.append(
                            {
                                "file": str(history_path),
                                "sheet": worksheet.title,
                                "row": excel_row,
                                "languages": added,
                            }
                        )
                    if set(normalized_languages).issubset(match.translations):
                        unresolved.remove(source_text)
        finally:
            workbook.close()

    results = [
        {
            "source_text": query,
            "status": matches[query].status(normalized_languages),
            "translations": matches[query].translations,
            "evidence": matches[query].evidence,
        }
        for query in normalized_queries
    ]
    return {
        "mode": "exact_only",
        "languages": list(normalized_languages),
        "query_count": len(normalized_queries),
        "exact_count": sum(item["status"] == "exact" for item in results),
        "partial_count": sum(item["status"] == "partial" for item in results),
        "miss_count": sum(item["status"] == "miss" for item in results),
        "scanned_history_files": scanned_files,
        "results": results,
    }
