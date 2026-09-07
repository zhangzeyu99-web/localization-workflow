"""Cross-row UI checks and explicit-history reuse warnings; never rewrite text."""
from __future__ import annotations

import re
import unicodedata
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from utils.quality_harness_terms import _detect_columns, _is_glossary_sheet, _is_support_sheet


GROUP_HARD_ISSUES = {"ui_series_inconsistency", "series_number_mismatch"}
ROMAN = re.compile(r"^(.*?[\u3400-\u9fff].*?)([IVXLCDM]+)$")
LEVEL = re.compile(r"(?<!\d)(\d{1,4})级")


def normalized(text):
    return re.sub(r"\s+", " ", unicodedata.normalize("NFKC", str(text or ""))).strip()


def _issue(row, kind, message, severity="error", **evidence):
    return {**row, "check_type": kind, "severity": severity, "message": message,
            "auto_fix": "", **evidence}


def group_issues(rows):
    groups = defaultdict(list)
    for row in rows:
        source = normalized(row["source"])
        target = normalized(row["translation"])
        if len(source) > 60 or re.search(r"[<>{}\[\]\n]", row["source"]):
            continue
        roman = ROMAN.fullmatch(source)
        levels = list(LEVEL.finditer(source))
        if roman:
            stem, number = roman.groups()
            key = (row.get("file"), row.get("sheet"), "roman", stem.strip())
            target_number = re.search(r"(?<![A-Za-z])([IVXLCDM]+)$", target)
            valid = bool(target_number and target_number[1] == number)
            signature = target[:target_number.start()] + "#" if target_number else target
        elif len(levels) == 1 and re.search(r"及以上|及以下|以上|以下|至少|至多", source):
            match = levels[0]
            number = match[1]
            stem = source[:match.start()] + "#级" + source[match.end():]
            key = (row.get("file"), row.get("sheet"), "level", stem)
            pattern = rf"(?<!\d){re.escape(number)}(?!\d)"
            valid = bool(re.search(pattern, target))
            signature = re.sub(pattern, "#", target)
        else:
            continue
        groups[key].append((row, number, signature, valid))

    issues = []
    for key, group in groups.items():
        if len({item[1] for item in group}) < 2:
            continue
        invalid = [item for item in group if not item[3]]
        if invalid:
            issues.extend(_issue(row, "series_number_mismatch", "Series member has a missing or different ordinal/level")
                          for row, _, _, _ in invalid)
        elif len({item[2] for item in group}) > 1:
            # Report every member; majority spelling is not evidence of correctness.
            issues.extend(_issue(row, "ui_series_inconsistency", "Equivalent source series uses different target stems or formats",
                                 group_source=key[-1], target_signature=signature)
                          for row, _, signature, _ in group)
    return issues


def history_issues(rows, history_rows):
    by_target = defaultdict(list)
    for old in history_rows:
        by_target[normalized(old["translation"])].append(old)
    issues = []
    for row in rows:
        target = normalized(row["translation"])
        # Common short labels legitimately translate several source phrases identically.
        if len(target) < 12:
            continue
        matches = [old for old in by_target.get(target, [])
                   if normalized(old["source"]) != normalized(row["source"])]
        if matches:
            evidence = [{k: old.get(k) for k in ("file", "sheet", "row", "source")} for old in matches[:5]]
            issues.append(_issue(row, "source_drift_tm_conflict",
                                 "Translation also belongs to a different historical source; review meaning before reuse",
                                 severity="warning", history_evidence=evidence))
    return issues


def load_history_rows(paths, lang, wanted_targets):
    result = []
    for raw_path in paths:
        path = Path(raw_path)
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for ws in wb.worksheets:
                if ws.max_row == 1 and ws.max_column == 1:
                    ws.reset_dimensions()
                if _is_glossary_sheet(ws) or _is_support_sheet(ws):
                    continue
                _, source_col, target_col = _detect_columns(ws, lang=lang)
                if source_col is None or target_col is None:
                    continue
                max_col = max(source_col, target_col) + 1
                for index, values in enumerate(ws.iter_rows(min_row=2, max_col=max_col, values_only=True), 2):
                    source, target = values[source_col], values[target_col]
                    if isinstance(source, str) and isinstance(target, str) and normalized(target) in wanted_targets:
                        result.append(dict(file=str(path), sheet=ws.title, row=index,
                                           source=source, translation=target))
        finally:
            wb.close()
    return result
