"""Local competitor-reference cache and review-table workflow.

The module deliberately keeps retrieval, language judgement, and workbook authoring
separate.  It never translates text and never treats a fuzzy candidate as an
approved equivalent.
"""
from __future__ import annotations

import csv
import hashlib
import json
import re
import shutil
import subprocess
import time
import zipfile
from collections import Counter
from collections import defaultdict
from difflib import SequenceMatcher
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries

from .large_text_multilingual_gate import TOKEN_RE, is_auto_protected_token


CACHE_SCHEMA = "competitor-reference-cache/v1"
MANIFEST_SCHEMA = "reference-suggestion-manifest/v1"
REVIEW_SCHEMA = "reference-suggestion-review/v1"
_PLACEHOLDER_RE = re.compile(
    r"<@\d+>|\{\{[^{}]+\}\}|\{[^{}]+\}|\$\{[^{}]+\}|%\{[^{}]+\}|%(?:\d+\$)?[a-zA-Z]|##\d+##|##\d+"
)
# Chinese letters must not hide adjacent source numbers (for example, 获得10个).
# Keep ASCII identifiers such as level10 protected from number-only matching.
_NUMBER_RE = re.compile(r"(?<![A-Za-z0-9_])\d+(?:[.,]\d+)*(?:%|[KMB])?(?![A-Za-z0-9_])", re.IGNORECASE)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _sha256_json(value: object) -> str:
    raw = json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


def _read_json(path: Path) -> object:
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _write_json(path: Path, value: object) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _read_records(path: Path) -> list[dict[str, object]]:
    path = Path(path)
    text = path.read_text(encoding="utf-8")
    if path.suffix.lower() == ".jsonl":
        return [json.loads(line) for line in text.splitlines() if line.strip()]
    value = json.loads(text)
    if isinstance(value, dict) and isinstance(value.get("records"), list):
        value = value["records"]
    if not isinstance(value, list) or not all(isinstance(row, dict) for row in value):
        raise ValueError("review file must contain a JSON array, records array, or JSONL objects")
    return value


def _write_jsonl(path: Path, rows: Iterable[dict[str, object]]) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="\n") as stream:
        for row in rows:
            stream.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")


def _text(value: object) -> str:
    return ("" if value is None else str(value)).replace("\r\n", "\n").replace("\r", "\n").strip()


def _raw_text(value: object) -> str:
    """Preserve source text exactly enough for provenance and length accounting."""
    return "" if value is None else str(value)


def _normalize_query(value: object) -> str:
    return re.sub(r"\s+", " ", _text(value)).casefold()


def _query_grams(value: str) -> set[str]:
    compact = re.sub(r"\s+", "", value)
    return {compact[index : index + 2] for index in range(max(0, len(compact) - 1))}


def _source_snapshot(path: Path, *, version: str = "") -> dict[str, object]:
    path = Path(path).resolve()
    if not path.is_file():
        raise FileNotFoundError(path)
    return {"path": str(path), "sha256": sha256_file(path), "version": version}


def _verify_snapshots(snapshots: Iterable[dict[str, object]], label: str) -> None:
    errors: list[str] = []
    for item in snapshots:
        path = Path(str(item.get("path") or ""))
        expected = str(item.get("sha256") or "")
        if not path.is_file():
            errors.append(f"missing {path}")
            continue
        actual = sha256_file(path)
        if actual != expected:
            errors.append(f"drift {path}: expected {expected}, got {actual}")
    if errors:
        raise ValueError(f"{label} source drift blocks reuse: " + "; ".join(errors))


def _assert_safe_outputs(outputs: Iterable[Path], inputs: Iterable[Path], label: str) -> None:
    resolved_outputs = [Path(path).resolve() for path in outputs]
    resolved_inputs = {Path(path).resolve() for path in inputs}
    if len(set(resolved_outputs)) != len(resolved_outputs):
        raise ValueError(f"{label} outputs must be distinct")
    collisions = [str(path) for path in resolved_outputs if path in resolved_inputs]
    if collisions:
        raise ValueError(f"{label} output collides with an input: {', '.join(collisions)}")


def _sheet_specs(source: dict[str, object]) -> list[dict[str, object]]:
    specs = source.get("sheets")
    if specs is None:
        specs = [{"name": source.get("sheet"), "header_row": source.get("header_row", 1), "columns": source.get("columns")}]
    if not isinstance(specs, list) or not specs:
        raise ValueError("each competitor source requires sheet or sheets")
    return [dict(item) for item in specs]


def _header_map(row: Iterable[object]) -> dict[str, int]:
    result: dict[str, int] = {}
    for index, value in enumerate(row, start=1):
        header = _text(value)
        if header and header not in result:
            result[header] = index
    return result


def _iter_competitor_xlsx(source: dict[str, object], snapshot: dict[str, object]) -> Iterator[dict[str, object]]:
    workbook = load_workbook(Path(str(snapshot["path"])), read_only=True, data_only=False)
    try:
        for sheet_spec in _sheet_specs(source):
            sheet_name = str(sheet_spec.get("name") or "")
            if sheet_name not in workbook.sheetnames:
                raise ValueError(f"competitor sheet not found: {sheet_name}")
            sheet = workbook[sheet_name]
            header_row = int(sheet_spec.get("header_row") or 1)
            headers = _header_map(next(sheet.iter_rows(min_row=header_row, max_row=header_row, values_only=True)))
            columns = dict(sheet_spec.get("columns") or {})
            languages = dict(columns.get("languages") or {})
            required = {"module": columns.get("module"), "key": columns.get("key"), **languages}
            missing = [str(header) for header in required.values() if str(header) not in headers]
            if missing:
                raise ValueError(f"missing competitor columns in {sheet_name}: {', '.join(missing)}")
            max_col = max(headers[str(header)] for header in required.values())
            for row_number, row in enumerate(
                sheet.iter_rows(min_row=header_row + 1, min_col=1, max_col=max_col, values_only=True),
                start=header_row + 1,
            ):
                texts = {str(code): _raw_text(row[headers[str(header)] - 1]) for code, header in languages.items()}
                key = _text(row[headers[str(columns["key"])] - 1])
                module = _text(row[headers[str(columns["module"])] - 1])
                if not key and not any(texts.values()):
                    continue
                yield {
                    "id": f"{source['id']}::{sheet_name}::{row_number}",
                    "source_id": str(source["id"]),
                    "sheet": sheet_name,
                    "physical_row": row_number,
                    "module": module,
                    "key": key,
                    "texts": texts,
                }
    finally:
        workbook.close()


def _iter_competitor_csv(source: dict[str, object], snapshot: dict[str, object]) -> Iterator[dict[str, object]]:
    columns = dict(source.get("columns") or {})
    languages = dict(columns.get("languages") or {})
    with Path(str(snapshot["path"])).open("r", encoding=str(source.get("encoding") or "utf-8-sig"), newline="") as stream:
        reader = csv.DictReader(stream)
        missing = [str(v) for v in [columns.get("module"), columns.get("key"), *languages.values()] if str(v) not in (reader.fieldnames or [])]
        if missing:
            raise ValueError("missing competitor CSV columns: " + ", ".join(missing))
        for row_number, row in enumerate(reader, start=2):
            texts = {str(code): _raw_text(row.get(str(header))) for code, header in languages.items()}
            key = _text(row.get(str(columns.get("key"))))
            if not key and not any(texts.values()):
                continue
            yield {
                "id": f"{source['id']}::CSV::{row_number}", "source_id": str(source["id"]),
                "sheet": "CSV", "physical_row": row_number,
                "module": _text(row.get(str(columns.get("module")))), "key": key,
                "texts": texts,
            }


def build_reference_cache(config_path: Path, output_path: Path) -> dict[str, object]:
    started = time.perf_counter()
    config = _read_json(config_path)
    if not isinstance(config, dict) or not isinstance(config.get("sources"), list) or not config["sources"]:
        raise ValueError("cache config requires a non-empty sources list")
    _assert_safe_outputs([output_path], [config_path, *(Path(str(row.get("path") or "")) for row in config["sources"])], "cache build")
    snapshots: list[dict[str, object]] = []
    records: list[dict[str, object]] = []
    seen_ids: set[str] = set()
    for raw in config["sources"]:
        source = dict(raw)
        source_id = str(source.get("id") or "").strip()
        if not source_id or source_id in seen_ids:
            raise ValueError("competitor source ids must be non-empty and unique")
        seen_ids.add(source_id)
        snapshot = _source_snapshot(Path(str(source.get("path") or "")), version=str(source.get("version") or ""))
        snapshot["id"] = source_id
        snapshots.append(snapshot)
        iterator = _iter_competitor_csv(source, snapshot) if Path(str(snapshot["path"])).suffix.lower() == ".csv" else _iter_competitor_xlsx(source, snapshot)
        records.extend(iterator)
    cache = {
        "schema": CACHE_SCHEMA, "config_path": str(Path(config_path).resolve()),
        "config_sha256": sha256_file(config_path), "sources": snapshots, "records": records,
        "record_count": len(records), "processing_seconds": round(time.perf_counter() - started, 6),
    }
    cache["content_sha256"] = _sha256_json({"sources": snapshots, "records": records})
    _write_json(output_path, cache)
    return cache


def load_reference_cache(cache_path: Path, *, verify_sources: bool = True) -> dict[str, object]:
    cache = _read_json(cache_path)
    if not isinstance(cache, dict) or cache.get("schema") != CACHE_SCHEMA:
        raise ValueError("unsupported competitor cache schema")
    if cache.get("content_sha256") != _sha256_json({"sources": cache.get("sources"), "records": cache.get("records")}):
        raise ValueError("competitor cache content hash mismatch")
    if verify_sources:
        config_path = Path(str(cache.get("config_path") or ""))
        if not config_path.is_file() or sha256_file(config_path) != cache.get("config_sha256"):
            raise ValueError("competitor catalog config drift blocks cache reuse")
        _verify_snapshots(cache.get("sources") or [], "competitor cache")
    return cache


def query_reference_cache(
    cache_path: Path, queries: Iterable[str], *, module_hint: str = "", top_k: int = 3, similarity_threshold: float = 0.72,
) -> dict[str, object]:
    started = time.perf_counter()
    if top_k < 1 or top_k > 3:
        raise ValueError("top_k must be between 1 and 3")
    cache = load_reference_cache(cache_path, verify_sources=True)
    sources_by_id = {str(row["id"]): row for row in cache["sources"]}
    search_values: list[list[str]] = []
    exact_index: dict[str, list[int]] = defaultdict(list)
    gram_index: dict[str, set[int]] = defaultdict(set)
    for index, record in enumerate(cache["records"]):
        texts = dict(record.get("texts") or {})
        values = list(dict.fromkeys(_normalize_query(texts.get(code, "")) for code in ("zh_cn", "zh-cn", "zh", "cn", "en") if _text(texts.get(code, ""))))
        search_values.append(values)
        for value in values:
            exact_index[value].append(index)
            for gram in _query_grams(value):
                gram_index[gram].add(index)
    output: list[dict[str, object]] = []
    hint = _normalize_query(module_hint)
    for query in queries:
        normalized = _normalize_query(query)
        if not normalized:
            continue
        exact: list[dict[str, object]] = []
        for index in exact_index.get(normalized, []):
            record = cache["records"][index]
            exact.append({**sources_by_id[str(record["source_id"])], **record, "retrieval_relation": "exact_text", "similarity": 1.0})
        exact.sort(key=lambda row: (0 if hint and hint in _normalize_query(row.get("module")) else 1, str(row["id"])))
        hits = exact[:top_k]
        if len(hits) < top_k:
            similar: list[tuple[float, dict[str, object]]] = []
            exact_ids = {str(row["id"]) for row in exact}
            candidate_indexes: set[int] = set()
            for gram in _query_grams(normalized):
                candidate_indexes.update(gram_index.get(gram, set()))
            for index in candidate_indexes:
                record = cache["records"][index]
                if str(record["id"]) in exact_ids:
                    continue
                values = search_values[index]
                plausible = [value for value in values if min(len(normalized), len(value)) / max(len(normalized), len(value), 1) >= similarity_threshold * 0.75]
                if not plausible:
                    continue
                module_bonus = 0.03 if hint and hint in _normalize_query(record.get("module")) else 0.0
                score = max(SequenceMatcher(None, normalized, value).ratio() for value in plausible) + module_bonus
                if score >= similarity_threshold:
                    similar.append((min(score, 0.999999), {**sources_by_id[str(record["source_id"])], **record}))
            similar.sort(key=lambda pair: (-pair[0], str(pair[1]["id"])))
            hits.extend({**record, "retrieval_relation": "similar_candidate", "similarity": round(score, 6)} for score, record in similar[: top_k - len(hits)])
        output.append({
            "query": query,
            "status": "exact_text" if exact else ("similar_candidates" if hits else "no_reliable_match"),
            "warning": "similar candidates are context only; they are not semantic equivalents or approved translations",
            "hits": hits,
        })
    return {
        "schema": "competitor-reference-query/v1", "cache_path": str(Path(cache_path).resolve()),
        "cache_sha256": sha256_file(cache_path), "queries": output,
        "elapsed_seconds": round(time.perf_counter() - started, 6),
    }


def _col(value: object) -> int:
    if isinstance(value, int):
        return value
    return column_index_from_string(str(value).strip())


def _support_snapshots(config: dict[str, object]) -> list[dict[str, object]]:
    values: list[dict[str, object]] = []
    for label in ("glossary", "context"):
        raw = config.get(label) or []
        paths = raw if isinstance(raw, list) else [raw]
        for path in paths:
            snapshot = _source_snapshot(Path(str(path)))
            snapshot["role"] = label
            values.append(snapshot)
    return values


def _candidate_rows(sheet, spec: dict[str, object], source_hash: str, review_context_signature: str) -> tuple[list[dict[str, object]], dict[str, object]]:
    _, _, sheet_max_column, sheet_max_row = range_boundaries(sheet.calculate_dimension(force=True))
    columns = dict(spec.get("columns") or {})
    target_map: dict[str, int] = {}
    for lang, column in dict(columns.get("targets") or {}).items():
        normalized_lang = str(lang).upper()
        if normalized_lang in target_map:
            raise ValueError(f"duplicate target language after normalization: {normalized_lang}")
        target_map[normalized_lang] = _col(column)
    if not target_map:
        raise ValueError("each task sheet requires target columns")
    key_col, source_col = _col(columns.get("key")), _col(columns.get("source"))
    note_col = _col(columns["note"]) if columns.get("note") else None
    role_columns = [key_col, source_col, *target_map.values(), *([note_col] if note_col else [])]
    if len(set(role_columns)) != len(role_columns):
        raise ValueError("key, source, target, and note columns must be distinct")
    if any(column < 1 or column > sheet_max_column for column in role_columns):
        raise ValueError("configured columns must exist inside the source sheet")
    has_header = spec.get("has_header")
    if not isinstance(has_header, bool):
        raise ValueError("has_header must be explicit true or false")
    first_data_row = 2 if has_header else 1
    matrix = list(sheet.iter_rows(min_row=1, max_row=sheet_max_row, min_col=1, max_col=sheet_max_column))
    at = lambda row, column: matrix[row - 1][column - 1]
    formulas = [cell.coordinate for row in matrix for cell in row if cell.data_type == "f"]
    styles: list[dict[str, object]] = []
    for row_number in range(1, sheet_max_row + 1):
        for column in range(1, sheet_max_column + 1):
            cell = at(row_number, column)
            font = getattr(cell, "font", None)
            alignment = getattr(cell, "alignment", None)
            styles.append({
                "row": row_number, "column": column, "number_format": getattr(cell, "number_format", None),
                "font": {"name": getattr(font, "name", None), "size": getattr(font, "sz", None), "bold": bool(getattr(font, "bold", False)), "italic": bool(getattr(font, "italic", False))},
                "alignment": {"horizontal": getattr(alignment, "horizontal", None), "vertical": getattr(alignment, "vertical", None), "wrap_text": getattr(alignment, "wrap_text", None)},
            })
    candidates: list[dict[str, object]] = []
    for row_number in range(first_data_row, sheet_max_row + 1):
        key = _raw_text(at(row_number, key_col).value)
        source = _raw_text(at(row_number, source_col).value)
        originals = {lang: _raw_text(at(row_number, col).value) for lang, col in target_map.items()}
        if not _text(key) and not _text(source) and not any(_text(value) for value in originals.values()):
            continue
        if not _text(key) or not _text(source):
            raise ValueError(f"active row {sheet.title}!{row_number} requires key and source text")
        for lang, original in originals.items():
            if not _text(original):
                raise ValueError(f"review target is blank: {sheet.title}!{get_column_letter(target_map[lang])}{row_number}")
            record_id = f"{sheet.title}!R{row_number}!{lang}"
            binding = {
                "source_sha256": source_hash, "review_context_signature": review_context_signature,
                "sheet": sheet.title, "row": row_number, "lang": lang, "key": key,
                "source": source, "original": original,
            }
            candidates.append({
                "id": record_id, "sheet": sheet.title, "row": row_number, "lang": lang, "key": key,
                "source": source, "original": original, "original_length": len(original),
                "decision": "", "suggested": "", "reason": "", "reference_relation": "",
                "evidence_ids": [], "basis": "", "evaluation_model": "", "controller_model": "",
                "controller_status": "", "batch_signature": _sha256_json(binding),
            })
    if formulas:
        raise ValueError("formula-bearing review cells are not supported in v1: " + ", ".join(formulas[:10]))
    return candidates, {
        "name": sheet.title, "has_header": has_header, "physical_rows": sheet_max_row,
        "physical_columns": sheet_max_column, "first_data_row": first_data_row,
        "columns": {"key": key_col, "source": source_col, "targets": target_map, "note": note_col},
        "styles": styles,
    }


def _unsupported_xlsx_features(path: Path) -> list[str]:
    unsupported: set[str] = set()
    with zipfile.ZipFile(path) as archive:
        if any(name.startswith("xl/tables/") for name in archive.namelist()):
            unsupported.add("structured tables")
        for name in archive.namelist():
            if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                xml = archive.read(name)
                for marker, label in (
                    (b"<mergeCell", "merged cells"),
                    (b"<conditionalFormatting", "conditional formatting"),
                    (b"<dataValidations", "data validation"),
                ):
                    if marker in xml:
                        unsupported.add(label)
    return sorted(unsupported)


def prepare_review(config_path: Path, work_dir: Path) -> dict[str, object]:
    started = time.perf_counter()
    config = _read_json(config_path)
    if not isinstance(config, dict):
        raise ValueError("task config must be a JSON object")
    source = Path(str(config.get("input") or "")).resolve()
    source_snapshot = _source_snapshot(source)
    unsupported = _unsupported_xlsx_features(source)
    if unsupported:
        raise ValueError("fixed-layout v1 does not support: " + ", ".join(unsupported))
    cache_path = Path(str(config.get("reference_cache") or "")).resolve()
    cache = load_reference_cache(cache_path, verify_sources=True)
    support = _support_snapshots(config)
    review_context_signature = _sha256_json({
        "task_config_sha256": sha256_file(config_path),
        "source": source_snapshot,
        "support_sources": support,
        "reference_cache": {
            "content_sha256": cache["content_sha256"],
            "config_path": cache["config_path"],
            "config_sha256": cache["config_sha256"],
            "sources": cache["sources"],
        },
    })
    workbook = load_workbook(source, read_only=True, data_only=False)
    source_sheetnames = list(workbook.sheetnames)
    candidates: list[dict[str, object]] = []
    sheet_manifests: list[dict[str, object]] = []
    try:
        formula_cells: list[str] = []
        for workbook_sheet in workbook.worksheets:
            _, _, max_column, max_row = range_boundaries(workbook_sheet.calculate_dimension(force=True))
            formula_cells.extend(
                f"{workbook_sheet.title}!{cell.coordinate}"
                for row in workbook_sheet.iter_rows(min_row=1, max_row=max_row, min_col=1, max_col=max_column)
                for cell in row
                if cell.data_type == "f"
            )
        if formula_cells:
            raise ValueError("formula-bearing workbooks are not supported in v1: " + ", ".join(formula_cells[:10]))
        specs = config.get("sheets")
        if not isinstance(specs, list) or not specs:
            raise ValueError("task config requires a non-empty sheets list")
        seen: set[str] = set()
        for raw in specs:
            spec = dict(raw)
            sheet_name = str(spec.get("name") or "")
            if not sheet_name or sheet_name in seen or sheet_name not in workbook.sheetnames:
                raise ValueError(f"task sheet must be present and unique: {sheet_name}")
            seen.add(sheet_name)
            rows, sheet_manifest = _candidate_rows(
                workbook[sheet_name], spec, str(source_snapshot["sha256"]), review_context_signature,
            )
            candidates.extend(rows)
            sheet_manifests.append(sheet_manifest)
    finally:
        workbook.close()
    if not candidates:
        raise ValueError("task contains no actual target cells to review")
    if len({row["id"] for row in candidates}) != len(candidates):
        raise ValueError("candidate ids are not unique")
    work_dir = Path(work_dir).resolve()
    candidates_path = work_dir / "review_candidates.jsonl"
    manifest_path = work_dir / "manifest.json"
    _assert_safe_outputs(
        [candidates_path, manifest_path],
        [config_path, source, cache_path, *(Path(str(row["path"])) for row in support), *(Path(str(row["path"])) for row in cache["sources"])],
        "prepare",
    )
    _write_jsonl(candidates_path, candidates)
    manifest = {
        "schema": MANIFEST_SCHEMA, "config_path": str(Path(config_path).resolve()), "config_sha256": sha256_file(config_path),
        "source": source_snapshot, "source_sheetnames": source_sheetnames, "sheets": sheet_manifests,
        "review_context_signature": review_context_signature,
        "support_sources": support, "reference_cache": {"path": str(cache_path), "sha256": sha256_file(cache_path), "content_sha256": cache["content_sha256"], "sources": cache["sources"]},
        "candidates": {"path": str(candidates_path), "sha256": sha256_file(candidates_path), "count": len(candidates)},
        "elapsed_seconds": round(time.perf_counter() - started, 6),
    }
    manifest["batch_id"] = _sha256_json({"source": source_snapshot, "sheets": sheet_manifests, "support": support, "reference": manifest["reference_cache"], "candidates_sha256": manifest["candidates"]["sha256"]})
    _write_json(manifest_path, manifest)
    return {**manifest, "manifest_path": str(manifest_path)}


def load_manifest(manifest_path: Path, *, verify_fresh: bool = True) -> dict[str, object]:
    manifest = _read_json(manifest_path)
    if not isinstance(manifest, dict) or manifest.get("schema") != MANIFEST_SCHEMA:
        raise ValueError("unsupported review manifest schema")
    if verify_fresh:
        config_path = Path(str(manifest.get("config_path") or ""))
        if not config_path.is_file() or sha256_file(config_path) != manifest.get("config_sha256"):
            raise ValueError("task config drift blocks review reuse")
        _verify_snapshots([manifest["source"], *(manifest.get("support_sources") or [])], "task")
        cache_info = dict(manifest["reference_cache"])
        cache_path = Path(str(cache_info["path"]))
        if sha256_file(cache_path) != cache_info["sha256"]:
            raise ValueError("reference cache file drift blocks review reuse")
        cache = load_reference_cache(cache_path, verify_sources=True)
        if cache["content_sha256"] != cache_info["content_sha256"]:
            raise ValueError("reference cache content drift blocks review reuse")
        candidates = Path(str(manifest["candidates"]["path"]))
        if sha256_file(candidates) != manifest["candidates"]["sha256"]:
            raise ValueError("candidate file drift blocks review reuse")
    return manifest


def _protected(text: str) -> Counter[str]:
    shared = [token for token in TOKEN_RE.findall(text) if is_auto_protected_token(token)]
    return Counter([*shared, *_PLACEHOLDER_RE.findall(text), *_NUMBER_RE.findall(text)])


def _fill_rgb(cell) -> str | None:
    color = getattr(getattr(cell, "fill", None), "fgColor", None)
    if getattr(color, "type", None) == "rgb" and getattr(color, "rgb", None):
        return str(color.rgb)[-6:].upper()
    return None


def _style_sig(cell) -> tuple[object, object, object, object]:
    font = getattr(cell, "font", None)
    protection = getattr(cell, "protection", None)
    return (getattr(font, "name", None), getattr(font, "sz", None), getattr(cell, "number_format", None), getattr(protection, "locked", None))


def _sheet_dimensions(sheet) -> tuple[int, int]:
    dimension = sheet.calculate_dimension(force=True)
    _, _, max_column, max_row = range_boundaries(dimension)
    return max_row, max_column


def _normalize_decision(value: object) -> str:
    normalized = _text(value).upper()
    return {"保持原样": "KEEP", "建议修改": "FIX"}.get(normalized, normalized)


def import_review(manifest_path: Path, review_path: Path, output_path: Path, report_path: Path) -> dict[str, object]:
    started = time.perf_counter()
    manifest = load_manifest(manifest_path, verify_fresh=True)
    _assert_safe_outputs(
        [output_path, report_path],
        [manifest_path, review_path, Path(str(manifest["source"]["path"])), Path(str(manifest["candidates"]["path"])), Path(str(manifest["reference_cache"]["path"])), Path(str(manifest["config_path"])), *(Path(str(row["path"])) for row in manifest.get("support_sources") or []), *(Path(str(row["path"])) for row in manifest["reference_cache"].get("sources") or [])],
        "review import",
    )
    candidates = _read_records(Path(str(manifest["candidates"]["path"])))
    reviews = _read_records(review_path)
    expected = {str(row["id"]): row for row in candidates}
    actual_ids = [str(row.get("id") or "") for row in reviews]
    duplicates = sorted(key for key, count in Counter(actual_ids).items() if key and count > 1)
    missing = sorted(set(expected) - set(actual_ids))
    extra = sorted(set(actual_ids) - set(expected))
    errors: list[dict[str, object]] = []
    if duplicates:
        errors.append({"code": "duplicate_review_id", "ids": duplicates})
    if missing:
        errors.append({"code": "missing_review_id", "ids": missing})
    if extra:
        errors.append({"code": "extra_review_id", "ids": extra})
    cache = load_reference_cache(Path(str(manifest["reference_cache"]["path"])), verify_sources=True)
    evidence = {str(row["id"]): row for row in cache["records"]}
    accepted: list[dict[str, object]] = []
    for review in reviews:
        record_id = str(review.get("id") or "")
        candidate = expected.get(record_id)
        if candidate is None:
            continue
        for field in ("sheet", "row", "lang", "key", "source", "original", "batch_signature"):
            if review.get(field) != candidate.get(field):
                errors.append({"code": "review_binding_mismatch", "id": record_id, "field": field})
        decision = _normalize_decision(review.get("decision"))
        suggested = _raw_text(review.get("suggested"))
        reason = _text(review.get("reason"))
        relation = _text(review.get("reference_relation"))
        basis = _text(review.get("basis"))
        evidence_ids = review.get("evidence_ids") or []
        if isinstance(evidence_ids, str):
            evidence_ids = [item.strip() for item in evidence_ids.split(",") if item.strip()]
        if decision not in {"KEEP", "FIX"}:
            errors.append({"code": "invalid_decision", "id": record_id})
        if decision == "KEEP" and suggested != "":
            errors.append({"code": "keep_suggestion_must_be_blank", "id": record_id})
        if decision == "FIX" and not _text(suggested):
            errors.append({"code": "fix_suggestion_required", "id": record_id})
        if decision == "FIX" and suggested == candidate["original"]:
            errors.append({"code": "fix_must_change_translation", "id": record_id})
        if not reason:
            errors.append({"code": "reason_required", "id": record_id})
        if relation not in {"corresponding", "similar", "no_match"}:
            errors.append({"code": "invalid_reference_relation", "id": record_id})
        if basis not in {"reference", "independent"}:
            errors.append({"code": "invalid_basis", "id": record_id})
        unknown_evidence = sorted(set(str(item) for item in evidence_ids) - set(evidence))
        if unknown_evidence:
            errors.append({"code": "unknown_evidence_id", "id": record_id, "evidence_ids": unknown_evidence})
        if (basis == "reference" or relation == "corresponding") and not evidence_ids:
            errors.append({"code": "reference_requires_evidence", "id": record_id})
        if relation == "no_match" and evidence_ids:
            errors.append({"code": "no_match_cannot_cite_evidence", "id": record_id})
        if not _text(review.get("evaluation_model")) or not _text(review.get("controller_model")) or _text(review.get("controller_status")).upper() != "ACCEPTED":
            errors.append({"code": "two_stage_review_required", "id": record_id})
        final_text = str(candidate["original"]) if decision == "KEEP" else suggested
        required_tokens = _protected(str(candidate["source"])) | _protected(str(candidate["original"]))
        final_tokens = _protected(final_text)
        missing_tokens = list((required_tokens - final_tokens).elements())
        if missing_tokens:
            errors.append({"code": "protected_content_missing", "id": record_id, "tokens": missing_tokens})
        extra_tokens = [token for token in (final_tokens - required_tokens).elements() if not _NUMBER_RE.fullmatch(token)]
        if extra_tokens:
            errors.append({"code": "protected_content_added", "id": record_id, "tokens": extra_tokens})
        green_eligible = False
        if decision == "KEEP" and relation == "corresponding" and evidence_ids:
            lang = str(candidate["lang"])
            for evidence_id in evidence_ids:
                texts = dict(evidence.get(str(evidence_id), {}).get("texts") or {})
                values = [value for code, value in texts.items() if str(code).upper().replace("-", "_") == lang.replace("-", "_")]
                green_eligible = green_eligible or any(_raw_text(value) == str(candidate["original"]) for value in values)
        accepted.append({
            **candidate, "decision": decision, "suggested": suggested, "final": final_text,
            "reason": reason, "reference_relation": relation, "evidence_ids": [str(item) for item in evidence_ids],
            "basis": basis, "evaluation_model": _text(review.get("evaluation_model")),
            "controller_model": _text(review.get("controller_model")), "controller_status": "ACCEPTED",
            "original_length": len(str(candidate["original"])), "final_length": len(final_text),
            "length_delta": len(final_text) - len(str(candidate["original"])), "green_eligible": green_eligible,
        })
    report = {
        "schema": "reference-suggestion-import-report/v1", "status": "blocked" if errors else "accepted",
        "batch_id": manifest["batch_id"], "candidate_count": len(candidates), "review_count": len(reviews),
        "error_count": len(errors), "errors": errors, "elapsed_seconds": round(time.perf_counter() - started, 6),
    }
    _write_json(report_path, report)
    if errors:
        raise ValueError(f"review import blocked with {len(errors)} error(s); report={report_path}")
    envelope = {
        "schema": REVIEW_SCHEMA, "status": "accepted", "batch_id": manifest["batch_id"],
        "manifest_path": str(Path(manifest_path).resolve()), "manifest_sha256": sha256_file(manifest_path),
        "records": accepted, "record_count": len(accepted), "import_report": str(Path(report_path).resolve()),
    }
    envelope["content_sha256"] = _sha256_json(accepted)
    _write_json(output_path, envelope)
    return {**report, "accepted_review": str(Path(output_path).resolve()), "content_sha256": envelope["content_sha256"]}


def load_accepted_review(path: Path, manifest: dict[str, object]) -> dict[str, object]:
    value = _read_json(path)
    if not isinstance(value, dict) or value.get("schema") != REVIEW_SCHEMA or value.get("status") != "accepted":
        raise ValueError("export requires an accepted review envelope")
    if value.get("batch_id") != manifest.get("batch_id") or value.get("content_sha256") != _sha256_json(value.get("records")):
        raise ValueError("accepted review does not match this batch")
    if value.get("manifest_sha256") != sha256_file(Path(str(value.get("manifest_path")))):
        raise ValueError("accepted review manifest drift")
    return value


def _runtime_paths() -> tuple[Path, Path, Path]:
    root = Path.home() / ".cache" / "codex-runtimes" / "codex-primary-runtime"
    node = root / "dependencies" / "node" / "bin" / "node.exe"
    modules = root / "dependencies" / "node" / "node_modules"
    marker = root / "plugins" / "openai-primary-runtime" / "plugins" / "spreadsheets" / "skills" / "spreadsheets" / "container_tools" / "mark_artifact_operation_started.mjs"
    if not node.is_file() or not modules.is_dir() or not marker.is_file():
        raise FileNotFoundError("bundled spreadsheet runtime not found; run load_workspace_dependencies")
    return node, modules, marker


def export_review(manifest_path: Path, accepted_path: Path, output_path: Path, work_dir: Path, builder_path: Path, *, mark_operation: bool = False) -> dict[str, object]:
    started = time.perf_counter()
    manifest = load_manifest(manifest_path, verify_fresh=True)
    review = load_accepted_review(accepted_path, manifest)
    _assert_safe_outputs(
        [output_path],
        [manifest_path, accepted_path, builder_path, Path(str(manifest["source"]["path"])), Path(str(manifest["candidates"]["path"])), Path(str(manifest["reference_cache"]["path"])), Path(str(manifest["config_path"])), *(Path(str(row["path"])) for row in manifest.get("support_sources") or []), *(Path(str(row["path"])) for row in manifest["reference_cache"].get("sources") or [])],
        "export",
    )
    work_dir = Path(work_dir).resolve()
    work_dir.mkdir(parents=True, exist_ok=True)
    node, modules, marker = _runtime_paths()
    local_builder = work_dir / "reference_suggestion_export.mjs"
    shutil.copyfile(builder_path, local_builder)
    module_link = work_dir / "node_modules"
    if not module_link.exists():
        subprocess.run(["cmd", "/c", "mklink", "/J", str(module_link), str(modules)], check=True, capture_output=True, text=True)
    if mark_operation:
        subprocess.run([str(node), str(marker), "--operation-kind", "edit", "--expected-output-count", "1", "--output-format", "xlsx"], check=True)
    spec = {
        "manifest": manifest, "review": review, "output": str(Path(output_path).resolve()),
        "preview_short": str(work_dir / "preview_short.png"), "preview_long": str(work_dir / "preview_long.png"),
        "verification": str(work_dir / "artifact_verification.ndjson"),
    }
    spec_path = work_dir / "export_spec.json"
    _write_json(spec_path, spec)
    completed = subprocess.run([str(node), str(local_builder), str(spec_path)], check=True, capture_output=True, text=True, encoding="utf-8")
    payload = json.loads(completed.stdout.strip().splitlines()[-1])
    return {**payload, "elapsed_seconds": round(time.perf_counter() - started, 6)}


def verify_export(manifest_path: Path, accepted_path: Path, output_path: Path, report_path: Path) -> dict[str, object]:
    started = time.perf_counter()
    manifest = load_manifest(manifest_path, verify_fresh=True)
    review = load_accepted_review(accepted_path, manifest)
    _assert_safe_outputs(
        [report_path],
        [manifest_path, accepted_path, output_path, Path(str(manifest["source"]["path"])), Path(str(manifest["candidates"]["path"])), Path(str(manifest["reference_cache"]["path"])), Path(str(manifest["config_path"])), *(Path(str(row["path"])) for row in manifest.get("support_sources") or []), *(Path(str(row["path"])) for row in manifest["reference_cache"].get("sources") or [])],
        "verify",
    )
    source_path = Path(str(manifest["source"]["path"]))
    output_path = Path(output_path)
    if not output_path.is_file():
        raise FileNotFoundError(output_path)
    source = load_workbook(source_path, read_only=True, data_only=False)
    output = load_workbook(output_path, read_only=True, data_only=False)
    errors: list[dict[str, object]] = []
    try:
        if source.sheetnames != output.sheetnames:
            errors.append({"code": "sheetnames_changed", "source": source.sheetnames, "output": output.sheetnames})
        records_by_sheet: dict[str, list[dict[str, object]]] = {}
        for record in review["records"]:
            records_by_sheet.setdefault(str(record["sheet"]), []).append(record)
        specs_by_name = {str(spec["name"]): spec for spec in manifest["sheets"]}
        for name in source.sheetnames:
            if name in specs_by_name:
                continue
            src, out = source[name], output[name]
            src_rows, src_columns = _sheet_dimensions(src)
            out_rows, out_columns = _sheet_dimensions(out)
            if (src_rows, src_columns) != (out_rows, out_columns):
                errors.append({"code": "unrequested_sheet_dimensions_changed", "sheet": name})
                continue
            for row in range(1, src_rows + 1):
                for col in range(1, src_columns + 1):
                    src_cell, out_cell = src.cell(row, col), out.cell(row, col)
                    if src_cell.value != out_cell.value or src_cell.data_type != out_cell.data_type:
                        errors.append({"code": "unrequested_sheet_cell_changed", "sheet": name, "cell": f"{get_column_letter(col)}{row}"})
                    if _style_sig(src_cell) != _style_sig(out_cell):
                        errors.append({"code": "unrequested_sheet_style_changed", "sheet": name, "cell": f"{get_column_letter(col)}{row}"})
        for spec in manifest["sheets"]:
            name = str(spec["name"])
            src, out = source[name], output[name]
            target_cols = {str(lang): int(col) for lang, col in dict(spec["columns"]["targets"]).items()}
            inserted_before = lambda col: sum(1 for target in target_cols.values() if target < col)
            expected_columns = int(spec["physical_columns"]) + len(target_cols) + (0 if spec["columns"].get("note") else 1)
            out_rows, out_columns = _sheet_dimensions(out)
            if out_rows != int(spec["physical_rows"]) or out_columns != expected_columns:
                errors.append({"code": "target_sheet_dimensions_changed", "sheet": name, "expected": [spec["physical_rows"], expected_columns], "actual": [out_rows, out_columns]})
            note_col = spec["columns"].get("note")
            note_output_col = int(note_col) + inserted_before(int(note_col)) if note_col else expected_columns
            for row in range(1, int(spec["physical_rows"]) + 1):
                for col in range(1, int(spec["physical_columns"]) + 1):
                    mapped = col + inserted_before(col)
                    src_cell, out_cell = src.cell(row, col), out.cell(row, mapped)
                    if col != note_col and (src_cell.value != out_cell.value or src_cell.data_type != out_cell.data_type):
                        errors.append({"code": "source_cell_changed", "sheet": name, "source": f"{get_column_letter(col)}{row}", "output": f"{get_column_letter(mapped)}{row}"})
                    if col != note_col and _style_sig(src_cell) != _style_sig(out_cell):
                        errors.append({"code": "source_style_changed", "sheet": name, "source": f"{get_column_letter(col)}{row}", "output": f"{get_column_letter(mapped)}{row}"})
                if note_col:
                    prior = _raw_text(src.cell(row, int(note_col)).value)
                    actual_note = _raw_text(out.cell(row, note_output_col).value)
                    has_records = any(int(item["row"]) == row for item in records_by_sheet.get(name, []))
                    if has_records and prior and not actual_note.startswith(prior + "\n"):
                        errors.append({"code": "prior_note_not_preserved", "sheet": name, "row": row})
                    if not has_records and actual_note != prior:
                        errors.append({"code": "blank_row_note_changed", "sheet": name, "row": row})
            for record in records_by_sheet.get(name, []):
                row = int(record["row"])
                original_col = target_cols[str(record["lang"])] + inserted_before(target_cols[str(record["lang"])] )
                suggestion_col = original_col + 1
                suggestion = _text(out.cell(row, suggestion_col).value)
                expected = str(record["suggested"]) if record["decision"] == "FIX" else ""
                if suggestion != expected:
                    errors.append({"code": "suggestion_mismatch", "id": record["id"], "expected": expected, "actual": suggestion})
                note = _text(out.cell(row, note_output_col).value)
                relation_label = {"corresponding": "对应", "similar": "近似", "no_match": "无可靠命中"}[str(record["reference_relation"])]
                basis_label = "参考竞品" if record["basis"] == "reference" else "项目判断"
                evidence_label = ",".join(record["evidence_ids"]) if record["evidence_ids"] else "无可靠竞品对应"
                decision_label = "保持原样" if record["decision"] == "KEEP" else "建议修改"
                expected_note = f"{record['lang']} {decision_label} {record['original_length']}→{record['final_length']}：{record['reason']}；{relation_label}；{basis_label}；{evidence_label}"
                if note.splitlines().count(expected_note) != 1:
                    errors.append({"code": "note_content_mismatch", "id": record["id"]})
                if out.cell(row, suggestion_col).alignment.wrap_text is not True or out.cell(row, suggestion_col).alignment.vertical != "center":
                    errors.append({"code": "suggestion_readability", "id": record["id"]})
                if out.cell(row, note_output_col).alignment.wrap_text is not True:
                    errors.append({"code": "note_readability", "id": record["id"]})
                if record["decision"] == "FIX":
                    if _fill_rgb(out.cell(row, original_col)) != "FFF2CC":
                        errors.append({"code": "fix_original_not_yellow", "id": record["id"]})
                    if _fill_rgb(out.cell(row, suggestion_col)) != "DDEBF7":
                        errors.append({"code": "fix_suggestion_not_blue", "id": record["id"]})
                elif record["green_eligible"] and _fill_rgb(out.cell(row, original_col)) != "E2F0D9":
                    errors.append({"code": "corresponding_keep_not_green", "id": record["id"]})
    finally:
        source.close()
        output.close()
    report = {
        "schema": "reference-suggestion-verify-report/v1", "status": "blocked" if errors else "verified",
        "batch_id": manifest["batch_id"], "error_count": len(errors), "errors": errors,
        "source_sha256": manifest["source"]["sha256"], "output_sha256": sha256_file(output_path),
        "elapsed_seconds": round(time.perf_counter() - started, 6),
    }
    _write_json(report_path, report)
    if errors:
        raise ValueError(f"export verification blocked with {len(errors)} error(s); report={report_path}")
    return report
