import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook
from utils.large_text_multilingual_gate import cache_lint, readback_gate, apply_dry_run
from utils.large_text_multilingual_proofread import run_deep_proofread, _validate_audit, _validate_suggestions
from utils.large_text_multilingual_runner import build_manifest


class AcceptanceIntegrityTests(unittest.TestCase):
    def test_cache_and_actual_readback_reject_same_semantic_and_token_errors(self):
        for source,target,kind in [
            ('编队2英雄攻击敌方基地时攻击力提升','El ATQ del héroe aumenta al defender la base','semantic_entity_number_missing'),
            ('{0}攻击，{0}防守','{0} ataca y defiende','protected_token_missing'),
            ('获得{0}金币','Obtén {0} monedas y {1} gemas','protected_token_extra'),
        ]:
            with self.subTest(kind=kind), tempfile.TemporaryDirectory() as tmp:
                root=Path(tmp);cache=root/'cache.jsonl'
                cache.write_text(json.dumps(dict(key='1',cn=source,translations={'ES':target}),ensure_ascii=False),encoding='utf-8')
                self.assertIn(kind,cache_lint(cache,target_langs=['ES'])['hard_by_type'])
                delivery=root/'delivery';delivery.mkdir();w=Workbook();s=w.active
                s.append(['ID','中文','西班牙语']);s.append(['1',source,target]);w.save(delivery/'final.xlsx');w.close()
                report=readback_gate(delivery,target_langs=['ES'])
                self.assertIn(kind,report['hard_by_type']);self.assertFalse(report['readback_verified'])
                self.assertEqual(report['checked_target_cells'],1)

    def test_empty_cache_directory_and_unknown_source_cannot_pass(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);cache=root/'empty.jsonl';cache.write_text('',encoding='utf-8')
            self.assertIn('empty_cache',cache_lint(cache,target_langs=['EN'])['hard_by_type'])
            delivery=root/'delivery';delivery.mkdir()
            self.assertIn('empty_delivery',readback_gate(delivery,target_langs=['EN'])['hard_by_type'])
            w=Workbook();s=w.active;s.append(['ID','EN']);s.append(['1','Go']);w.save(delivery/'bad.xlsx');w.close()
            r=readback_gate(delivery,target_langs=['EN']);self.assertIn('source_column_missing',r['hard_by_type'])

    def test_dry_run_cannot_overwrite_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'source.xlsx';w=Workbook();w.save(p);w.close();before=p.read_bytes()
            with self.assertRaises(ValueError):apply_dry_run(p,p)
            self.assertEqual(p.read_bytes(),before)

    def test_audit_unresolved_and_changed_keep_rejected(self):
        fix=dict(review_key='1',lang='ES',cn='攻击',current='Ataca',suggested='Defiende')
        with self.assertRaises(ValueError):_validate_audit([fix],[dict(review_key='1',lang='ES',decision='ACCEPT',final='Defiende',unresolved_issues=['wrong action'])])
        with self.assertRaises(ValueError):_validate_suggestions([dict(review_key='1',translations={'ES':'Ataca'})],[dict(review_key='1',lang='ES',status='KEEP',suggested='Defiende')],['ES'])
        with self.assertRaises(ValueError):_validate_audit([{**fix,'status':'KEEP'}],[dict(review_key='1',lang='ES',decision='REVERT',final='Ataca',reason='do not change')])

    def test_scalar_numbers_tags_and_real_newlines(self):
        from utils.large_text_multilingual_gate import pair_integrity_issues
        for src,dst,lang,kind in [
            ('攻击2个目标','Ataca objetivos','ES','number_missing'),
            ('持续2回合','Dura turnos','ES','number_missing'),
            ('攻击2个敌人','Ataca inimigos dos aliados','PT','number_missing'),
            ('#G4#{0}#n#伤害','{0} de daño','ES','protected_token_missing'),
            ('开始\n结束','Start End','EN','newline_mismatch'),
        ]:
            with self.subTest(src=src):self.assertIn(kind,{k for k,d in pair_integrity_issues({'cn':src},lang,dst)})
        for src,dst,lang in [
            ('攻击2个目标','Ataca a dos objetivos','ES'),
            ('每发出1辆货车','Para cada caminhão enviado','PT'),
            ('获得宝石*10,5分钟加速*2,50k金币*2','Get Gems*10, 5-Min Speedup*2, 50k Coins*2','EN'),
        ]:
            self.assertNotIn('number_missing',{k for k,d in pair_integrity_issues({'cn':src},lang,dst)})

    def test_keep_is_independently_audited_and_can_be_corrected(self):
        seen=[]
        class Reviewer:
            def review_batch(self,rows,target_langs):
                return [dict(review_key=r['review_key'],lang='ES',status='KEEP',suggested=r['translations']['ES'],reason='incorrectly kept') for r in rows]
        class Auditor:
            def audit_batch(self,rows):
                seen.extend(rows)
                return [dict(review_key=r['review_key'],lang='ES',decision='REVISE',final='Ataca a un enemigo',reason='restore action and target') for r in rows]
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);items=root/'items.jsonl';initial=root/'initial.jsonl'
            payload=json.dumps(dict(key='1',cn='攻击1个敌人',translations={'ES':'Defiende'}),ensure_ascii=False)+'\n'
            items.write_text(payload,encoding='utf-8');initial.write_text(payload,encoding='utf-8')
            m=build_manifest(work_dir=root/'work',items_jsonl=items,source_rows_jsonl=None,target_langs=['ES'],workbook_count=1,relay_config=None,proofread_mode='full')
            result=run_deep_proofread(Path(m['manifest_path']),initial_cache=initial,reviewer=Reviewer(),auditor=Auditor())
            self.assertEqual(len(seen),1);self.assertEqual(seen[0]['status'],'KEEP')
            self.assertEqual(result.suggested_changes,0);self.assertEqual(result.changed_cells,1)
            self.assertEqual(json.loads(result.final_cache.read_text(encoding='utf-8'))['translations']['ES'],'Ataca a un enemigo')
            self.assertEqual(json.loads(result.summary_json.read_text(encoding='utf-8'))['audited_keep_cells'],1)
