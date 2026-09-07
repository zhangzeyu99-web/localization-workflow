import json
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook

from utils.quality_harness import scan_workbook
from utils.review_consistency import group_issues, history_issues


def rows(pairs, sheet="UI"):
    return [dict(file="sample.xlsx", sheet=sheet, row=i + 2, id=i + 1,
                 source=source, translation=target) for i, (source, target) in enumerate(pairs)]


def test_shared_regression_fixtures():
    fixture = json.loads((Path(__file__).parents[1] / "fixtures/quality_regression.json").read_text(encoding="utf-8"))
    for case in fixture["review_consistency_cases"]:
        items = rows(case["rows"])
        issues = group_issues(items)
        if "history" in case:
            issues += history_issues(items, rows(case["history"]))
        assert sorted({x["check_type"] for x in issues}) == sorted(case["expected"]), case["name"]


def test_groups_do_not_cross_sheets():
    items = rows([("挑战9级及以上怪物", "Defeat Lv. 9+ monsters")])
    items += rows([("挑战11级及以上怪物", "Defeat monsters of level 11 or above")], "Tutorial")
    assert group_issues(items) == []


def test_duplicate_rows_are_not_a_series():
    assert group_issues(rows([("礼包Ⅰ", "Gift I"), ("礼包Ⅰ", "Pack I")])) == []


def test_consistent_but_incorrect_ordinal_is_blocked():
    issues = group_issues(rows([("礼包Ⅰ", "Gift III"), ("礼包Ⅱ", "Gift IV")]))
    assert {x["check_type"] for x in issues} == {"series_number_mismatch"}


def test_equal_quantity_and_level_do_not_create_false_format_drift():
    assert group_issues(rows([
        ('击败9个9级以上怪物', 'Defeat 9 Lv. 9+ monsters'),
        ('击败9个11级以上怪物', 'Defeat 9 Lv. 11+ monsters'),
    ])) == []


def test_harness_entry_checks_series_and_explicit_history():
    with TemporaryDirectory() as tmp:
        current, history = Path(tmp) / "current.xlsx", Path(tmp) / "history.xlsx"
        for path, pairs in [(current, [("艾拉赠礼Ⅰ", "Lucky Gifts I"), ("艾拉赠礼Ⅱ", "Lucky Gift II")]),
                            (history, [("幸运赠礼Ⅰ", "Lucky Gifts I")])]:
            wb = Workbook()
            wb.active.append(["ID", "CN", "EN"])
            for i, pair in enumerate(pairs, 1):
                wb.active.append([i, *pair])
            wb.save(path)
            wb.close()
        result = scan_workbook(current, auto_discover_terms=False, history=[history])
        assert not result.passed
        assert result.issue_counts["ui_series_inconsistency"] == 2
        assert result.issue_counts["source_drift_tm_conflict"] == 1
        assert any(x["severity"] == "warning" for x in result.issues)


def test_name_category_in_notes_is_enforced_by_both_loaders():
    from utils.large_text_multilingual_pack import _load_terms
    with TemporaryDirectory() as tmp:
        terms, current = Path(tmp) / 'terms.xlsx', Path(tmp) / 'current.xlsx'
        wb = Workbook()
        wb.active.append(['ID', 'CN', 'EN', '备注'])
        wb.active.append([1, '艾拉', 'Aria', '英雄名'])
        wb.active.append([2, '试用', 'Trial', '角色确认后再改'])
        wb.save(terms)
        wb.close()
        wb = Workbook()
        wb.active.append(['ID', 'CN', 'EN'])
        wb.active.append([1, '艾拉赠礼Ⅰ', 'Lucky Gifts I'])
        wb.save(current)
        wb.close()
        result = scan_workbook(current, term_base=[terms])
        assert result.issue_counts['person_name_term_mismatch'] == 1
        loaded = {t['source']:t for t in _load_terms(terms, ['EN'])}
        assert loaded['艾拉']['required'] is True
        assert loaded['试用']['required'] is False


def test_note_category_fixture_and_named_glossary_sheet():
    from utils.quality_harness_terms import name_category_from_note, _collect_terms_from_workbook
    fixture = json.loads((Path(__file__).parents[1] / 'fixtures/quality_regression.json').read_text(encoding='utf-8'))
    for case in fixture['name_note_category_cases']:
        assert name_category_from_note(case['value']) == case['expected']
    wb = Workbook()
    wb.active.title = '术语'
    wb.active.append(['ID', 'CN', 'EN', 'ES', '备注'])
    wb.active.append([1, '艾拉', 'Aria', 'Aria', '英雄名'])
    received = []
    _collect_terms_from_workbook(wb, lambda *args: received.append(args))
    assert received[0][2] == '角色名'
    wb.close()


def test_explicit_category_is_not_overridden_by_notes():
    from utils.quality_harness_terms import _collect_terms_from_workbook
    wb = Workbook()
    wb.active.title = '术语'
    wb.active.append(['CN', 'EN', '分类', '备注'])
    wb.active.append(['艾拉', 'Aria', '通用词', '英雄名'])
    received = []
    _collect_terms_from_workbook(wb, lambda *args: received.append(args))
    assert received[0][2] == '通用词'
    wb.close()
