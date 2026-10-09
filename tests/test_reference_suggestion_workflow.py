from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from utils.reference_suggestion_workflow import (
    _runtime_paths,
    build_reference_cache,
    export_review,
    import_review,
    prepare_review,
    query_reference_cache,
    sha256_file,
    verify_export,
)


ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "reference_suggestion_source.xlsx"
BUILDER = ROOT / "scripts" / "reference_suggestion_export.mjs"


def _write_json(path: Path, value: object) -> None:
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def _make_batch(tmp_path: Path, *, reference_en: str = "Keep", source_en: str | None = None) -> dict[str, Path]:
    source = tmp_path / "source.xlsx"
    shutil.copyfile(FIXTURE, source)
    if source_en is not None:
        workbook = load_workbook(source)
        try:
            workbook["Sheet1"]["C3"] = source_en
            workbook.save(source)
        finally:
            workbook.close()
    reference = tmp_path / "reference.csv"
    with reference.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(["module", "key", "zh", "en", "es", "pt"])
        writer.writerow(["ui", "dup.key", "短词", reference_en, "Cambiar", "Mudar"])
        writer.writerow(["other", "dup.key", "手动", "Manual", "Manual", "Manual"])
        writer.writerow(["ui", "long.key", "获得10个<@1>奖励", "Get 10 <@1> rewards", "Obtén 10 <@1> premios", "Receba 10 <@1> recompensas"])
    catalog = tmp_path / "catalog.json"
    _write_json(catalog, {
        "schema": "competitor-reference-catalog/v1",
        "sources": [{
            "id": "fixture-v1", "version": "1", "path": str(reference),
            "columns": {"module": "module", "key": "key", "languages": {"zh_cn": "zh", "en": "en", "es": "es", "pt": "pt"}},
        }],
    })
    cache = tmp_path / "reference_cache.json"
    build_reference_cache(catalog, cache)
    context = tmp_path / "context.txt"
    context.write_text("stable fixture context\n", encoding="utf-8")
    task = tmp_path / "task.json"
    _write_json(task, {
        "input": str(source), "reference_cache": str(cache), "context": [str(context)],
        "sheets": [{
            "name": "Sheet1", "has_header": False,
            "columns": {"key": "A", "source": "B", "targets": {"EN": "C", "ES": "D", "PT": "E"}, "note": "F"},
        }],
    })
    work = tmp_path / "work"
    manifest = prepare_review(task, work)
    return {
        "source": source, "reference": reference, "catalog": catalog, "cache": cache,
        "context": context, "task": task, "work": work, "manifest": Path(manifest["manifest_path"]),
        "candidates": Path(manifest["candidates"]["path"]),
    }


def _read_jsonl(path: Path) -> list[dict[str, object]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def _reviews(batch: dict[str, Path]) -> list[dict[str, object]]:
    records = _read_jsonl(batch["candidates"])
    result: list[dict[str, object]] = []
    for item in records:
        review = {
            **item, "decision": "KEEP", "suggested": "", "reason": "fits the fixture context",
            "reference_relation": "no_match", "evidence_ids": [], "basis": "independent",
            "evaluation_model": "fixture-reviewer", "controller_model": "fixture-controller", "controller_status": "ACCEPTED",
        }
        if item["id"] == "Sheet1!R1!EN":
            review.update(reference_relation="corresponding", evidence_ids=["fixture-v1::CSV::2"], basis="reference")
        elif item["id"] == "Sheet1!R1!ES":
            review.update(decision="FIX", suggested="Cambiar mejor")
        elif item["id"] == "Sheet1!R3!ES":
            review.update(decision="FIX", suggested="10 <@1> premios")
        result.append(review)
    return result


def test_fixed_table_round_trip_preserves_rows_notes_styles_and_unrequested_sheet(tmp_path: Path) -> None:
    try:
        _runtime_paths()
    except FileNotFoundError:
        pytest.skip("Codex bundled Node/Artifact runtime is unavailable")
    batch = _make_batch(tmp_path)
    candidates = _read_jsonl(batch["candidates"])
    assert len(candidates) == 6
    assert candidates[0]["id"] == "Sheet1!R1!EN"
    assert {item["row"] for item in candidates} == {1, 3}
    assert json.loads(batch["manifest"].read_text(encoding="utf-8"))["sheets"][0]["physical_rows"] == 3

    source_hash = sha256_file(batch["source"])
    review_path = tmp_path / "review.json"
    accepted = tmp_path / "accepted.json"
    import_report = tmp_path / "import_report.json"
    _write_json(review_path, _reviews(batch))
    imported = import_review(batch["manifest"], review_path, accepted, import_report)
    assert imported["error_count"] == 0
    accepted_records = json.loads(accepted.read_text(encoding="utf-8"))["records"]
    green = next(item for item in accepted_records if item["id"] == "Sheet1!R1!EN")
    assert green["green_eligible"] is True

    output = tmp_path / "suggestions.xlsx"
    exported = export_review(batch["manifest"], accepted, output, batch["work"], BUILDER)
    assert exported["records"] == 6
    verified = verify_export(batch["manifest"], accepted, output, tmp_path / "verify.json")
    assert verified["status"] == "verified"
    assert sha256_file(batch["source"]) == source_hash

    workbook = load_workbook(output, read_only=True, data_only=False)
    try:
        sheet = workbook["Sheet1"]
        assert sheet.calculate_dimension(force=True) == "A1:I3"
        assert sheet["C1"].value == "Keep" and sheet["D1"].value is None
        assert sheet["E1"].value == "Cambiar" and sheet["F1"].value == "Cambiar mejor"
        assert sheet["G1"].value == "Mudar" and sheet["H1"].value is None
        assert sheet["G1"].number_format == "@"
        assert sheet["A2"].value is None and sheet["I2"].value is None
        note = sheet["I1"].value
        assert note.startswith("prior note\n")
        assert "EN 保持原样 4→4：fits the fixture context；对应；参考竞品；fixture-v1::CSV::2" in note
        assert "ES 建议修改 7→13：fits the fixture context；无可靠命中；项目判断；无可靠竞品对应" in note
        assert sheet["C1"].fill.fgColor.rgb[-6:] == "E2F0D9"
        assert sheet["E1"].fill.fgColor.rgb[-6:] == "FFF2CC"
        assert sheet["F1"].fill.fgColor.rgb[-6:] == "DDEBF7"
        assert sheet["F1"].alignment.wrap_text is True and sheet["F1"].alignment.vertical == "center"
        assert workbook["Untouched"]["B1"].value == 0
        assert workbook["Untouched"]["B2"].value == "stay"
    finally:
        workbook.close()


def test_review_binding_and_protected_content_are_hard_blockers(tmp_path: Path) -> None:
    batch = _make_batch(tmp_path)
    reviews = _reviews(batch)
    target = next(item for item in reviews if item["id"] == "Sheet1!R3!ES")
    target["suggested"] = "premios"
    review_path = tmp_path / "bad_review.json"
    report_path = tmp_path / "bad_report.json"
    _write_json(review_path, reviews)
    with pytest.raises(ValueError, match="review import blocked"):
        import_review(batch["manifest"], review_path, tmp_path / "accepted.json", report_path)
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert any(item["code"] == "protected_content_missing" for item in report["errors"])

    reviews = _reviews(batch)
    reviews[0]["original"] = "wrong batch text"
    _write_json(review_path, reviews)
    with pytest.raises(ValueError, match="review import blocked"):
        import_review(batch["manifest"], review_path, tmp_path / "accepted.json", report_path)
    report = json.loads(report_path.read_text(encoding="utf-8"))
    assert any(item["code"] == "review_binding_mismatch" and item["field"] == "original" for item in report["errors"])


def test_green_requires_the_same_untrimmed_translation(tmp_path: Path) -> None:
    batch = _make_batch(tmp_path, reference_en="Keep ")
    review_path = tmp_path / "review.json"
    accepted = tmp_path / "accepted.json"
    _write_json(review_path, _reviews(batch))
    import_review(batch["manifest"], review_path, accepted, tmp_path / "import_report.json")
    records = json.loads(accepted.read_text(encoding="utf-8"))["records"]
    item = next(row for row in records if row["id"] == "Sheet1!R1!EN")
    assert item["green_eligible"] is False


@pytest.mark.parametrize("suggested,accepted", [("Claim <@1> rewards", False), ("Claim 10 <@1> rewards", True)])
def test_chinese_adjacent_source_number_is_required_even_if_original_omits_it(tmp_path: Path, suggested: str, accepted: bool) -> None:
    batch = _make_batch(tmp_path, source_en="Get <@1> rewards")
    reviews = _reviews(batch)
    target = next(item for item in reviews if item["id"] == "Sheet1!R3!EN")
    target.update(decision="FIX", suggested=suggested)
    review_path = tmp_path / "review.json"
    report_path = tmp_path / "import_report.json"
    _write_json(review_path, reviews)
    if accepted:
        imported = import_review(batch["manifest"], review_path, tmp_path / "accepted.json", report_path)
        assert imported["status"] == "accepted"
    else:
        with pytest.raises(ValueError, match="review import blocked"):
            import_review(batch["manifest"], review_path, tmp_path / "accepted.json", report_path)
        report = json.loads(report_path.read_text(encoding="utf-8"))
        assert any(item["code"] == "protected_content_missing" and item["id"] == "Sheet1!R3!EN" and "10" in item["tokens"] for item in report["errors"])


def test_review_binding_tracks_context_and_reference_inputs(tmp_path: Path) -> None:
    batch = _make_batch(tmp_path)
    reviews = _reviews(batch)
    old_signatures = {str(item["batch_signature"]) for item in reviews}
    review_path = tmp_path / "old_review.json"
    _write_json(review_path, reviews)

    same = prepare_review(batch["task"], tmp_path / "same_context")
    same_candidates = _read_jsonl(Path(same["candidates"]["path"]))
    assert {str(item["batch_signature"]) for item in same_candidates} == old_signatures
    accepted = tmp_path / "same_accepted.json"
    imported = import_review(Path(same["manifest_path"]), review_path, accepted, tmp_path / "same_report.json")
    assert imported["status"] == "accepted"

    batch["context"].write_text("changed authoritative context\n", encoding="utf-8")
    changed_context = prepare_review(batch["task"], tmp_path / "changed_context")
    changed_candidates = _read_jsonl(Path(changed_context["candidates"]["path"]))
    assert {str(item["batch_signature"]) for item in changed_candidates}.isdisjoint(old_signatures)
    with pytest.raises(ValueError, match="review import blocked"):
        import_review(
            Path(changed_context["manifest_path"]), review_path,
            tmp_path / "context_accepted.json", tmp_path / "context_report.json",
        )
    report = json.loads((tmp_path / "context_report.json").read_text(encoding="utf-8"))
    assert any(item["code"] == "review_binding_mismatch" and item["field"] == "batch_signature" for item in report["errors"])

    batch["context"].write_text("stable fixture context\n", encoding="utf-8")
    reference_text = batch["reference"].read_text(encoding="utf-8")
    batch["reference"].write_text(reference_text + "ui,new.key,新词,New,New,Novo\n", encoding="utf-8")
    build_reference_cache(batch["catalog"], batch["cache"])
    changed_reference = prepare_review(batch["task"], tmp_path / "changed_reference")
    with pytest.raises(ValueError, match="review import blocked"):
        import_review(
            Path(changed_reference["manifest_path"]), review_path,
            tmp_path / "reference_accepted.json", tmp_path / "reference_report.json",
        )
    report = json.loads((tmp_path / "reference_report.json").read_text(encoding="utf-8"))
    assert any(item["code"] == "review_binding_mismatch" and item["field"] == "batch_signature" for item in report["errors"])


def test_cache_and_catalog_drift_block_reuse_and_queries_do_not_approve_semantics(tmp_path: Path) -> None:
    batch = _make_batch(tmp_path)
    result = query_reference_cache(batch["cache"], ["手册"], top_k=3, similarity_threshold=0.72)
    assert result["queries"][0]["status"] == "no_reliable_match"
    assert result["queries"][0]["hits"] == []

    catalog = json.loads(batch["catalog"].read_text(encoding="utf-8"))
    catalog["catalog_version"] = "2"
    _write_json(batch["catalog"], catalog)
    with pytest.raises(ValueError, match="catalog config drift"):
        query_reference_cache(batch["cache"], ["短词"])


@pytest.mark.parametrize(
    "targets,note,error",
    [
        ({"EN": "C", "en": "D"}, "F", "duplicate target language"),
        ({"EN": "A"}, "F", "must be distinct"),
        ({"EN": "C"}, "C", "must be distinct"),
    ],
)
def test_prepare_rejects_duplicate_or_overlapping_columns(tmp_path: Path, targets: dict[str, str], note: str, error: str) -> None:
    batch = _make_batch(tmp_path)
    config = json.loads(batch["task"].read_text(encoding="utf-8"))
    config["sheets"][0]["columns"]["targets"] = targets
    config["sheets"][0]["columns"]["note"] = note
    bad_config = tmp_path / "bad_task.json"
    _write_json(bad_config, config)
    with pytest.raises(ValueError, match=error):
        prepare_review(bad_config, tmp_path / "bad_work")
