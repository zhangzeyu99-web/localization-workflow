from __future__ import annotations

import json
import re
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook

from utils.large_text_multilingual_pack import _load_terms, prepare_pack, stable_row_key


LANGS = ["EN", "IDN", "DE", "FR", "ES", "PT", "RU", "IT", "TR", "TH"]


def create_workbook(path: Path, start_id: int, rows: int, unique_count: int = 145) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "UI" if "UI" in path.stem else "Language"
    sheet.append(["ID", "CN", *LANGS])
    for offset in range(rows):
        sheet.append([start_id + offset, f"文本{offset % unique_count}", *([None] * len(LANGS))])
    workbook.save(path)
    workbook.close()


def create_reference_workbook(path: Path, *, missing_reference: bool = False) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Language"
    sheet.append(["ID", "CN", "EN", "FR", "DE"])
    sheet.append([1, "烈焰斩", "Flame Strike", None, None])
    sheet.append([2, "太阳神殿", None if missing_reference else "Temple of the Sun", None, None])
    workbook.save(path)
    workbook.close()


class LargeTextMultilingualPackTests(unittest.TestCase):
    def test_declared_character_category_aliases_reach_cache_gate(self) -> None:
        from utils.large_text_multilingual_gate import cache_lint
        from utils.large_text_multilingual_pack import _term_hits

        fixture = json.loads((Path(__file__).parents[1] / 'fixtures' / 'quality_regression.json').read_text(encoding='utf-8'))
        for case in fixture['strict_term_category_cases']:
            with self.subTest(category=case['category']), tempfile.TemporaryDirectory() as tmp:
                root = Path(tmp)
                term_base = root / 'terms.xlsx'
                workbook = Workbook()
                workbook.active.append(['CN', 'EN', 'FR', '分类'])
                workbook.active.append(['艾拉', 'Ella', 'Ella', case['category']])
                workbook.save(term_base)
                workbook.close()
                terms = _load_terms(term_base, ['FR'])
                self.assertIs(terms[0]['strict'], case['required'])
                row = {'key': '1', 'cn': '你好，艾拉', 'source_mode': 'en',
                       'translation_source': 'Hello, Ella.',
                       'term_hits': _term_hits('你好，艾拉', terms),
                       'translations': {'FR': 'Bonjour, Emma.'}}
                cache = root / 'cache.jsonl'
                cache.write_text(json.dumps(row, ensure_ascii=False) + '\n', encoding='utf-8')
                report = cache_lint(cache, target_langs=['FR'], term_base=term_base)
                self.assertEqual(report['hard_by_type'], {'term_missing': 1} if case['required'] else {})
                row['translations']['FR'] = 'Bonjour, Ella.'
                cache.write_text(json.dumps(row, ensure_ascii=False) + '\n', encoding='utf-8')
                self.assertEqual(cache_lint(cache, target_langs=['FR'], term_base=term_base)['hard_blockers'], 0)

    def test_load_terms_ignores_stale_xlsx_dimension_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "terms.xlsx"
            rewritten = Path(tmp) / "rewritten.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["ID", "CN", "EN", "分类"])
            sheet.append([2179, "双生魔偶", "Clockwork Twins", "主角"])
            workbook.save(path)
            workbook.close()
            with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(rewritten, "w") as target:
                for info in source.infolist():
                    payload = source.read(info.filename)
                    if info.filename == "xl/worksheets/sheet1.xml":
                        payload = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', payload, count=1)
                    target.writestr(info, payload)
            terms = _load_terms(rewritten, ["EN"])

            self.assertEqual([term["source"] for term in terms], ["双生魔偶"])

    def test_english_primary_names_and_kinship_are_contextual(self) -> None:
        from utils.large_text_multilingual_gate import cache_lint
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            terms_path = root / 'terms.xlsx'
            workbook = Workbook()
            workbook.active.append(['CN', 'EN', 'FR', '分类'])
            workbook.active.append(['艾拉', 'Ella', 'Ella', '角色'])
            workbook.active.append(['姥姥', 'Grandma', 'Mamie', '角色'])
            workbook.active.append(['艾拉奶奶', 'Grandma Ella', 'Mamie Ella', '角色'])
            workbook.save(terms_path)
            workbook.close()
            terms = {term['source']: term for term in _load_terms(terms_path, ['FR'])}
            self.assertFalse(terms['姥姥']['strict'])
            self.assertTrue(terms['艾拉奶奶']['strict'])
            self.assertEqual(terms['艾拉']['reference_en'], 'Ella')
            for mode in ('en', 'cn+en'):
                source = root / f'{mode}.xlsx'
                workbook = Workbook()
                workbook.active.append(['ID', 'CN', 'EN', 'FR'])
                workbook.active.append([1, '艾拉说过了', 'You already said that.', None])
                workbook.active.append([2, '艾拉说过了', "Ella's already said that.", None])
                workbook.active.append([3, '看望姥姥', 'Visit Grandma.', None])
                workbook.save(source)
                workbook.close()
                pack = prepare_pack(inputs=[source], term_base=terms_path, history_dirs=[], target_langs=['FR'], work_dir=root/mode, source_mode=mode)
                rows = [json.loads(line) for line in pack.items_jsonl.read_text(encoding='utf-8').splitlines()]
                self.assertEqual(len(rows[0]['term_hits']), 0 if mode == 'en' else 1)
                self.assertEqual(len(rows[1]['term_hits']), 1)
                for row, value in zip(rows, ['Tu l’as déjà dit.', 'Ella l’a déjà dit.', 'Rendre visite à mamie.']):
                    row['translations'] = {'FR': value}
                cache = root / f'{mode}.jsonl'
                cache.write_text(''.join(json.dumps(row, ensure_ascii=False)+'\n' for row in rows), encoding='utf-8')
                report = cache_lint(cache, target_langs=['FR'], term_base=terms_path)
                self.assertEqual(report['hard_by_type'], {} if mode == 'en' else {'term_missing': 1})

    def test_load_terms_marks_main_character_names_as_required(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "terms.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["ID", "CN", "EN", "FR", "DE", "PT", "ES", "TR", "RU", "备注", "分类"])
            sheet.append(
                [
                    2179,
                    "双生魔偶",
                    "Clockwork Twins",
                    "Jumeaux Mécaniques",
                    "Uhrwerk-Zwillinge",
                    "Gêmeas Mecânicas",
                    "Gemelos Mecánicos",
                    "Kurmalı İkizler",
                    "Заводные близнецы",
                    None,
                    "主角",
                ]
            )
            workbook.save(path)
            workbook.close()

            terms = _load_terms(path, ["EN", "FR", "DE", "PT", "ES", "TR", "RU"])

            self.assertEqual(len(terms), 1)
            self.assertEqual(terms[0]["source"], "双生魔偶")
            self.assertEqual(terms[0]["category"], "主角")
            self.assertIs(terms[0]["required"], True)
            self.assertIs(terms[0]["strict"], True)

    def test_stable_row_key_includes_file_sheet_id_and_row(self) -> None:
        self.assertEqual(
            stable_row_key("a.xlsx", "Sheet1", 7, 1001),
            "a.xlsx::Sheet1::1001::7",
        )

    def test_prepare_pack_extracts_rows_and_deduplicates_unique_texts(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            first = root / "7.13新增.xlsx"
            second = root / "7.13UI新增.xlsx"
            create_workbook(first, 1000, 245)
            create_workbook(second, 2000, 22)

            result = prepare_pack(
                inputs=[first, second],
                term_base=None,
                history_dirs=[],
                target_langs=LANGS,
                work_dir=root / "work",
            )

            self.assertEqual(result.source_rows, 267)
            self.assertEqual(result.unique_items, 145)
            self.assertEqual(result.estimated_target_cells, 2670)
            self.assertLess(result.elapsed_seconds, 5.0)
            self.assertTrue(result.items_jsonl.exists())
            self.assertTrue(result.source_rows_jsonl.exists())
            rows = [json.loads(line) for line in result.items_jsonl.read_text(encoding="utf-8").splitlines()]
            self.assertEqual(len(rows), 267)
            self.assertEqual(rows[0]["context"], "language")
            self.assertEqual(rows[-1]["context"], "ui")
            self.assertEqual(set(rows[0]["translations"]), set())

    def test_prepare_pack_rejects_nonempty_target_cells(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "input.xlsx"
            create_workbook(path, 1, 1)
            from openpyxl import load_workbook

            workbook = load_workbook(path)
            workbook.active["C2"] = "Existing"
            workbook.save(path)

            with self.assertRaisesRegex(ValueError, "target column is not empty"):
                prepare_pack(
                    inputs=[path],
                    term_base=None,
                    history_dirs=[],
                    target_langs=LANGS,
                    work_dir=root / "work",
                )

    def test_prepare_pack_adds_english_reference_for_non_english_targets(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "input.xlsx"
            create_reference_workbook(path)

            result = prepare_pack(
                inputs=[path],
                term_base=None,
                history_dirs=[],
                target_langs=["FR", "DE"],
                work_dir=root / "work",
                source_mode="cn+en",
            )
            rows = [
                json.loads(line)
                for line in result.items_jsonl.read_text(encoding="utf-8").splitlines()
            ]
            stats = json.loads(result.prepare_stats.read_text(encoding="utf-8"))

            self.assertEqual(rows[0]["source_mode"], "cn+en")
            self.assertEqual(rows[0]["translation_source"], "烈焰斩")
            self.assertEqual(rows[0]["reference_en"], "Flame Strike")
            self.assertEqual(rows[0]["reference_en_status"], "usable")
            self.assertEqual(stats["english_reference_status"]["usable_rows"], 2)

    def test_prepare_pack_en_mode_rejects_incomplete_reference(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            path = root / "input.xlsx"
            create_reference_workbook(path, missing_reference=True)

            with self.assertRaisesRegex(ValueError, "complete usable English"):
                prepare_pack(
                    inputs=[path],
                    term_base=None,
                    history_dirs=[],
                    target_langs=["FR", "DE"],
                    work_dir=root / "work",
                    source_mode="en",
                )


if __name__ == "__main__":
    unittest.main()
