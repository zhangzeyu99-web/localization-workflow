import json
from pathlib import Path
from unittest.mock import Mock
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook

from utils.glossary_enrichment import accept_glossary, prepare_request, verify_receipt


def test_compact_review_skips_rejections_and_bounds_repeated_evidence():
    from utils.glossary_enrichment import compact_review_packet

    long_source = '说明' * 1500 + '星辉秘境' + '规则' * 1500
    evidence = {'file': 'source.xlsx', 'sheet': 'Sheet', 'row': 2, 'source': long_source}
    payload = {'candidates': [
        {'cn': '星辉秘境', 'decision': 'needs_context', 'category': '副本',
         'translations': {'EN': ''}, 'evidence': [evidence, evidence]},
        {'cn': '已领取', 'decision': 'reject', 'reason': 'low_value_prompt', 'evidence': [evidence]},
    ], 'existing_incomplete': [{'cn': '星门', 'missing_languages': ['ES'], 'translations': {'EN': 'Star Gate'},
                                'evidence': [{'file': 'source.xlsx', 'row': 3, 'source': '点击星门'}]}]}
    packet = compact_review_packet(payload)
    assert len(packet['candidates']) == 1
    assert packet['excluded_from_model_review'] == 1
    kept = packet['candidates'][0]['evidence']
    assert len(kept) == 1 and len(kept[0]['excerpt']) <= 240
    assert '星辉秘境' in kept[0]['excerpt'] and kept[0]['excerpt_only']
    assert kept[0]['row'] == 2 and 'source' not in kept[0]
    assert packet['existing_incomplete'][0]['missing_languages'] == ['ES']
    assert payload['candidates'][0]['evidence'][0]['source'] == long_source


def workbook(path, rows, headers=None):
    wb = Workbook()
    wb.active.append(headers or ['ID', 'CN', 'EN', 'ES', '分类'])
    for row in rows:
        wb.active.append(row)
    wb.save(path)
    wb.close()


def fixture(tmp_path):
    source, base, updated = (tmp_path / name for name in ('source.xlsx', 'base.xlsx', 'updated.xlsx'))
    workbook(source, [[1, '月门', None, None, None]])
    workbook(base, [[1, '星门', 'Star Gate', 'Puerta Estelar', '地名']])
    workbook(updated, [[1, '星门', 'Star Gate', 'Puerta Estelar', '地名'], [2, '月门', 'Moon Gate', 'Puerta Lunar', '地名']])
    candidates = tmp_path / 'candidates.json'
    candidates.write_text(json.dumps({'candidates': [{'cn': '月门', 'decision': 'include', 'category': '地名', 'evidence': [{'file': str(source), 'sheet': 'Sheet', 'row': 2, 'source': '月门'}]}]}), encoding='utf-8')
    request = tmp_path / 'request.json'
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'ES'], candidates=candidates, output=request)
    return source, base, updated, candidates, request


def test_accept_and_verify(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    receipt = tmp_path / 'receipt.json'
    result = accept_glossary(request=request, reviewed=candidates, updated=updated, output=receipt)
    assert result['approved_terms'] == 1
    assert verify_receipt(receipt, inputs=[source], term_base=updated, target_langs=['ES', 'EN'])['status'] == 'accepted'


def test_preflight_resolves_relative_paths_before_cross_repo_subprocess(tmp_path, monkeypatch):
    from scripts import run_glossary_preflight as cli

    source, base, _updated, candidates, _request = fixture(tmp_path)
    extractor = tmp_path / 'extractor' / 'scripts'
    extractor.mkdir(parents=True)
    (extractor / 'export_glossary_candidates.py').write_text('', encoding='utf-8')
    review = tmp_path / 'review.json'
    review.write_text('{}', encoding='utf-8')
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(cli.sys, 'argv', ['preflight', 'prepare', '--input', source.name,
        '--term-base', base.name, '--target-langs', 'EN,ES', '--glossary-workflow', 'extractor',
        '--work-dir', 'stage', '--review-file', review.name])

    def run(command, **kwargs):
        for option in ['--input', '--existing-glossary', '--output', '--review-file']:
            assert Path(command[command.index(option) + 1]).is_absolute()
        output = Path(command[command.index('--output') + 1])
        output.write_bytes(candidates.read_bytes())
        return SimpleNamespace(returncode=0, stdout='', stderr='')

    monkeypatch.setattr(cli.subprocess, 'run', run)
    assert cli.main() == 0
    assert (tmp_path / 'stage' / 'request.json').is_file()


def test_source_drift_blocks_before_pack_or_api(tmp_path):
    from utils.large_text_multilingual_pipeline import run_pipeline
    source, base, updated, candidates, request = fixture(tmp_path)
    receipt = tmp_path / 'receipt.json'
    accept_glossary(request=request, reviewed=candidates, updated=updated, output=receipt)
    workbook(source, [[1, '变更', None, None, None]])
    client = Mock()
    with pytest.raises(ValueError, match='drift'):
        run_pipeline(inputs=[source], term_base=updated, history_dirs=[], target_langs=['EN', 'ES'], task_dir=tmp_path / 'task', relay_config=None, proofread_mode='basic', translation_client=client, glossary_receipt=receipt)
    client.translate_batch.assert_not_called()


@pytest.mark.parametrize('problem', ['unresolved', 'missing_language', 'old_row_modified', 'omitted_candidate'])
def test_reject_incomplete_enrichment(tmp_path, problem):
    source, base, updated, candidates, request = fixture(tmp_path)
    review = json.loads(candidates.read_text(encoding='utf-8'))
    if problem == 'unresolved':
        review['candidates'][0]['decision'] = 'needs_context'
    elif problem == 'omitted_candidate':
        review['candidates'] = []
    elif problem == 'missing_language':
        workbook(updated, [[1, '星门', 'Star Gate', 'Puerta Estelar', '地名'], [2, '月门', 'Moon Gate', None, '地名']])
    else:
        workbook(updated, [[1, '星门', 'Changed', 'Puerta Estelar', '地名'], [2, '月门', 'Moon Gate', 'Puerta Lunar', '地名']])
    reviewed = tmp_path / 'reviewed.json'
    reviewed.write_text(json.dumps(review), encoding='utf-8')
    with pytest.raises(ValueError):
        accept_glossary(request=request, reviewed=reviewed, updated=updated, output=tmp_path / 'receipt.json')


def test_receipt_rejects_changed_language_or_glossary(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    receipt = tmp_path / 'receipt.json'
    accept_glossary(request=request, reviewed=candidates, updated=updated, output=receipt)
    with pytest.raises(ValueError, match='languages'):
        verify_receipt(receipt, inputs=[source], term_base=updated, target_langs=['EN'])
    workbook(updated, [])
    with pytest.raises(ValueError, match='drift'):
        verify_receipt(receipt, inputs=[source], term_base=updated, target_langs=['EN', 'ES'])


def test_existing_blank_translations_may_be_filled(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    workbook(base, [[1, '星门', 'Star Gate', None, '地名']])
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'ES'], candidates=candidates, output=request)
    result = accept_glossary(request=request, reviewed=candidates, updated=updated, output=tmp_path / 'receipt.json')
    assert result['filled_existing_cells'] == 1


def test_existing_hit_must_be_complete(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    workbook(base, [[1, '星门', 'Star Gate', None, '地名']])
    workbook(updated, [[1, '星门', 'Star Gate', None, '地名'], [2, '月门', 'Moon Gate', 'Puerta Lunar', '地名']])
    raw = json.loads(candidates.read_text(encoding='utf-8'))
    raw['existing_incomplete'] = [{'cn': '星门', 'missing_languages': ['ES']}]
    candidates.write_text(json.dumps(raw), encoding='utf-8')
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'ES'], candidates=candidates, output=request)
    with pytest.raises(ValueError, match='existing glossary translation'):
        accept_glossary(request=request, reviewed=candidates, updated=updated, output=tmp_path / 'receipt.json')


def test_appended_category_must_match_review(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    workbook(updated, [[1, '星门', 'Star Gate', 'Puerta Estelar', '地名'], [2, '月门', 'Moon Gate', 'Puerta Lunar', '道具']])
    with pytest.raises(ValueError, match='category'):
        accept_glossary(request=request, reviewed=candidates, updated=updated, output=tmp_path / 'receipt.json')


def test_indonesian_idn_header_takes_priority_over_identifier_id(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    for path, rows in [(base, [[1, '星门', 'Star Gate', 'Gerbang Bintang', '地名']]),
                       (updated, [[1, '星门', 'Star Gate', 'Gerbang Bintang', '地名'],
                                  [2, '月门', 'Moon Gate', 'Gerbang Bulan', '地名']])]:
        wb = Workbook()
        wb.active.append(['ID', 'CN', 'EN', 'IDN', '分类'])
        for row in rows:
            wb.active.append(row)
        wb.save(path)
        wb.close()
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'IDN'],
                    candidates=candidates, output=request)
    receipt = tmp_path / 'receipt.json'
    result = accept_glossary(request=request, reviewed=candidates, updated=updated, output=receipt)
    assert result['approved_terms'] == 1
    assert verify_receipt(receipt, inputs=[source], term_base=updated,
                          target_langs=['EN', 'IDN'])['status'] == 'accepted'


def test_indonesian_blank_translation_can_be_filled_without_changing_blank_id(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    headers = ['ID', 'CN', 'EN', 'IDN', '分类']
    workbook(base, [[None, '星门', 'Star Gate', None, '地名']], headers)
    workbook(updated, [[None, '星门', 'Star Gate', 'Gerbang Bintang', '地名'],
                       [2, '月门', 'Moon Gate', 'Gerbang Bulan', '地名']], headers)
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'IDN'],
                    candidates=candidates, output=request)
    result = accept_glossary(request=request, reviewed=candidates, updated=updated,
                            output=tmp_path / 'receipt.json')
    assert result['status'] == 'accepted'
    assert result['filled_existing_cells'] == 1


def test_indonesian_language_alias_cannot_authorize_filling_identifier_id(tmp_path):
    source, base, updated, candidates, request = fixture(tmp_path)
    headers = ['ID', 'CN', 'EN', 'IDN', '分类']
    workbook(base, [[None, '星门', 'Star Gate', 'Gerbang Bintang', '地名']], headers)
    workbook(updated, [['changed identifier', '星门', 'Star Gate', 'Gerbang Bintang', '地名'],
                       [2, '月门', 'Moon Gate', 'Gerbang Bulan', '地名']], headers)
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'IDN'],
                    candidates=candidates, output=request)
    with pytest.raises(ValueError, match='existing glossary rows changed'):
        accept_glossary(request=request, reviewed=candidates, updated=updated,
                        output=tmp_path / 'receipt.json')
    assert not (tmp_path / 'receipt.json').exists()


def test_accepted_glossary_receipt_passes_complete_pipeline(tmp_path):
    from utils.large_text_multilingual_pipeline import run_pipeline

    source, base, updated, candidates, request = fixture(tmp_path)
    workbook(source, [[1, '点击月门', None, None, None], [2, '立即领取', None, None, None]])
    prepare_request(inputs=[source], term_base=base, target_langs=['EN', 'ES'],
                    candidates=candidates, output=request)
    receipt = tmp_path / 'receipt.json'
    accept_glossary(request=request, reviewed=candidates, updated=updated, output=receipt)

    class Translator:
        checkpoint_identity = 'synthetic-glossary-review-translator'
        calls = 0

        def translate_batch(self, rows, target_langs):
            self.calls += 1
            translations = {'点击月门': {'EN': 'Open the Moon Gate', 'ES': 'Abre la Puerta Lunar'},
                            '立即领取': {'EN': 'Claim now', 'ES': 'Reclamar ahora'}}
            return [{'request_key': row['request_key'],
                     'translations': {lang: translations[row['cn']][lang] for lang in target_langs}}
                    for row in rows]

    translator = Translator()
    result = run_pipeline(inputs=[source], term_base=updated, history_dirs=[], target_langs=['EN', 'ES'],
                          task_dir=tmp_path / 'task', relay_config=None, proofread_mode='basic',
                          translation_client=translator, glossary_receipt=receipt)
    manifest = json.loads(result.manifest.read_text(encoding='utf-8'))
    assert manifest['status'] == 'complete'
    assert manifest['critical_path'][0] == 'glossary_enrichment'
    assert manifest['phase_status']['glossary_enrichment'] == 'done'
    assert manifest['artifacts']['glossary_receipt'] == str(receipt)
    assert manifest['phase_metrics']['glossary_enrichment'] == {'approved_terms': 1, 'rejected_terms': 0}
    assert translator.calls == 1 and result.hard_blockers == 0
    readback = json.loads((result.manifest.parent / 'readback_gate.json').read_text(encoding='utf-8'))
    assert readback['readback_verified'] and readback['hard_blockers'] == 0
    delivered = load_workbook(result.delivery_dir / source.name, read_only=False)
    original = load_workbook(source, read_only=False)
    try:
        assert [tuple(row[:4]) for row in delivered.active.values] == [
            ('ID', 'CN', 'EN', 'ES'),
            (1, '点击月门', 'Open the Moon Gate', 'Abre la Puerta Lunar'),
            (2, '立即领取', 'Claim now', 'Reclamar ahora'),
        ]
        assert [tuple(row[:4]) for row in original.active.values] == [
            ('ID', 'CN', 'EN', 'ES'), (1, '点击月门', None, None), (2, '立即领取', None, None),
        ]
    finally:
        delivered.close()
        original.close()
